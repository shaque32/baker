"""Synthetic case01: determinism, report/database agreement, draft key integrity, planted traps."""

import re
import sqlite3
from datetime import UTC, datetime, timedelta
from pathlib import Path

import openpyxl
import pytest

from core.contracts import GoldClaim
from eval.run_eval import load_jsonl
from eval.synthetic import case01_key as key
from eval.synthetic import case01_story as story
from eval.synthetic.generate import DEFAULT_OUT, affidavit_paragraphs, generate

COMMITTED = Path(__file__).resolve().parents[1] / DEFAULT_OUT
COMMITTED_FILES = (
    "case.json",  # pins the SHA-256 of item1.xlsx and item2.xlsx
    "affidavit_draft.md",
    "draft_gold.jsonl",
    "ANSWER_KEY_DRAFT.md",
)
RAW_TS = re.compile(r"^(\d+)/(\d+)/(\d{4}) (\d+):(\d\d):(\d\d) (AM|PM)\(UTC([+-])(\d+)\)$")


@pytest.fixture(scope="module")
def case(tmp_path_factory):
    out = tmp_path_factory.mktemp("case01")
    db = out / "case01.db"
    rendered = generate(out, db)
    conn = sqlite3.connect(db)
    yield out, conn, rendered
    conn.close()


@pytest.fixture(scope="module")
def gold(case) -> list[GoldClaim]:
    out, _, _ = case
    return load_jsonl(out / "draft_gold.jsonl", GoldClaim)


def ids_for(rendered, device: str, story_key: str) -> list[str]:
    r = next(x for x in rendered if x.device.key == device)
    return r.keys[story_key]


def parse_raw(raw: str) -> tuple[datetime, datetime]:
    """Printed time -> (naive printed wall clock, true UTC)."""
    m = RAW_TS.match(raw)
    assert m, raw
    mo, d, y, h, mi, s, ampm, sign, off = m.groups()
    hour = int(h) % 12 + (12 if ampm == "PM" else 0)
    wall = datetime(int(y), int(mo), int(d), hour, int(mi), int(s))
    offset = int(off) * (1 if sign == "+" else -1)
    return wall, (wall - timedelta(hours=offset)).replace(tzinfo=UTC)


# ---------------------------------------------------------------- determinism and size


@pytest.mark.parametrize("name", COMMITTED_FILES)
def test_committed_outputs_match_a_fresh_generation(case, name):
    out, _, _ = case
    assert (out / name).read_bytes() == (COMMITTED / name).read_bytes(), (
        f"{name} is stale: run `make synth` and commit the result"
    )


def test_about_two_thousand_messages_mixed_languages(case):
    _, conn, rendered = case
    total = conn.execute("SELECT count(*) FROM messages").fetchone()[0]
    ru = sum(m["lang"] == "ru" for r in rendered for m in r.messages)
    assert 1800 <= total <= 2200
    assert ru >= 200
    assert conn.execute("SELECT count(DISTINCT source_id) FROM messages").fetchone()[0] == 2


def test_foreign_keys_hold(case):
    _, conn, _ = case
    assert conn.execute("PRAGMA foreign_key_check").fetchall() == []


# ---------------------------------------------------------------- report <-> database


@pytest.mark.parametrize("source", ["item1", "item2"])
def test_database_rows_match_the_excel_report(case, source):
    out, conn, _ = case
    wb = openpyxl.load_workbook(out / f"{source}.xlsx", read_only=True)
    chats = {i: row for i, row in enumerate(wb["Chats"].iter_rows(values_only=True), start=1)}
    header = list(chats[2])
    body_col, ts_col = header.index("Body"), header.index("Timestamp")
    rows = conn.execute(
        "SELECT id, locator, body, ts_raw FROM messages WHERE source_id = ?", (source,)
    ).fetchall()
    assert len(rows) == len(chats) - 2
    for mid, locator, body, ts_raw in rows:
        assert mid == f"msg:{source}:{locator}"
        sheet, row_no = locator.split("!")
        row = chats[int(row_no)]
        assert sheet == "Chats"
        assert (row[body_col] or "") == body
        assert row[ts_col] == ts_raw
    wb.close()


def test_printed_times_round_trip_to_stored_utc(case):
    _, conn, _ = case
    for table in ("messages", "calls"):
        for ts_raw, ts_utc, off in conn.execute(
            f"SELECT ts_raw, ts_utc, ts_offset_min FROM {table}"
        ):
            wall, utc = parse_raw(ts_raw)
            assert utc.strftime("%Y-%m-%dT%H:%M:%SZ") == ts_utc
            assert (wall - utc.replace(tzinfo=None)) == timedelta(minutes=off)


def test_report_time_zones_differ_by_device(case):
    _, conn, _ = case
    offsets = dict(
        conn.execute(
            "SELECT source_id, group_concat(DISTINCT ts_offset_min) FROM messages GROUP BY 1"
        ).fetchall()
    )
    assert offsets["item1"] == "0"
    assert set(offsets["item2"].split(",")) == {"-300", "-240"}


# ---------------------------------------------------------------- draft answer key


def test_draft_key_shape(gold):
    assert len(gold) == 20
    assert len({g.claim_id for g in gold}) == 20
    split = {
        v: sum(g.gold_verdict == v for g in gold) for v in ("supported", "contradicted", "unproven")
    }
    single = [g for g in gold if not g.cross_device]
    single_split = {v: sum(g.gold_verdict == v for g in single) for v in split}
    assert sum(g.cross_device for g in gold) == 2
    assert single_split == {"supported": 8, "contradicted": 4, "unproven": 6}
    assert split == {"supported": 8, "contradicted": 6, "unproven": 6}


def test_draft_key_is_never_marked_final(case, gold):
    out, _, _ = case
    assert all(g.labeled_by == key.DRAFT_LABEL and "DRAFT" in g.labeled_by for g in gold)
    assert (
        (out / "ANSWER_KEY_DRAFT.md")
        .read_text(encoding="utf-8")
        .startswith("# case01 answer key: DRAFT")
    )


def test_every_cited_record_exists(case, gold):
    _, conn, _ = case
    for g in gold:
        assert g.key_evidence, g.claim_id
        for rid in g.key_evidence:
            hit = conn.execute(
                "SELECT 1 FROM messages WHERE id = ? UNION ALL SELECT 1 FROM calls WHERE id = ? "
                "UNION ALL SELECT 1 FROM contacts WHERE id = ? "
                "UNION ALL SELECT 1 FROM attachments WHERE id = ?",
                (rid, rid, rid, rid),
            ).fetchone()
            assert hit, f"{g.claim_id}: {rid} missing"


def test_claims_are_verbatim_in_their_affidavit_paragraph(case, gold):
    out, _, _ = case
    paras = {(p, n): text for p, n, text in affidavit_paragraphs()}
    md = (out / "affidavit_draft.md").read_text(encoding="utf-8")
    for g in gold:
        assert g.text in paras[(g.page, g.para_no)], g.claim_id
        assert f"{g.para_no}. {paras[(g.page, g.para_no)]}" in md


def test_quoted_text_in_claims_matches_cited_records_verbatim(case, gold):
    _, conn, _ = case
    for g in gold:
        cited = []
        for rid in g.key_evidence:
            row = conn.execute(
                "SELECT body FROM messages WHERE id = ? UNION ALL "
                "SELECT name FROM contacts WHERE id = ?",
                (rid, rid),
            ).fetchone()
            if row:
                cited.append(row[0])
        for quote in re.findall(r'"([^"]+)"', g.text):
            assert any(quote in c for c in cited), f"{g.claim_id}: {quote!r} not in cited records"


# ---------------------------------------------------------------- planted traps


def test_trap_gap_whatsapp_silent_but_other_apps_continue(case):
    _, conn, _ = case
    start, end = (
        story.WA_GAP_START.strftime("%Y-%m-%dT%H:%M:%SZ"),
        story.WA_GAP_END.strftime("%Y-%m-%dT%H:%M:%SZ"),
    )
    wa = conn.execute(
        "SELECT count(*) FROM messages m JOIN threads t ON t.id = m.thread_id "
        "WHERE t.app = 'WhatsApp' AND m.ts_utc >= ? AND m.ts_utc < ?",
        (start, end),
    ).fetchone()[0]
    other = conn.execute(
        "SELECT count(*) FROM messages m JOIN threads t ON t.id = m.thread_id "
        "WHERE m.source_id = 'item1' AND t.app != 'WhatsApp' AND m.ts_utc >= ? AND m.ts_utc < ? "
        "AND m.id IN (SELECT message_id FROM message_recipients r JOIN accounts a "
        "ON a.id = r.account_id WHERE a.identifier IN ('+12125550122', '7001002')"
        " UNION SELECT m2.id FROM messages m2 JOIN accounts a2 ON a2.id = m2.sender_account_id"
        " WHERE a2.identifier IN ('+12125550122', '7001002'))",
        (start, end),
    ).fetchone()[0]
    deleted_wa = conn.execute(
        "SELECT count(*) FROM messages m JOIN threads t ON t.id = m.thread_id "
        "WHERE t.app = 'WhatsApp' AND m.deleted_flag = 1"
    ).fetchone()[0]
    assert wa == 0
    assert other >= 4
    assert deleted_wa == 0


def test_trap_handle_change_same_user_id(case):
    out, _, _ = case
    wb = openpyxl.load_workbook(out / "item1.xlsx", read_only=True)
    header, *rows = list(wb["Chats"].iter_rows(min_row=2, values_only=True))
    wb.close()
    col = {name: i for i, name in enumerate(header)}
    froms = {
        r[col["From"]]
        for r in rows
        if r[col["Identifier"]] == "5551234" and r[col["Direction"]] == "Incoming"
    }
    assert froms == {"5551234 @alex92", "5551234 @northstar"}


def test_trap_count_claim_matches_data(case):
    _, conn, _ = case
    involved = (
        "SELECT m.id FROM messages m JOIN accounts a ON a.id = m.sender_account_id "
        "WHERE a.identifier = '5551234' UNION SELECT r.message_id FROM message_recipients r "
        "JOIN accounts a ON a.id = r.account_id WHERE a.identifier = '5551234'"
    )
    n = conn.execute(
        "SELECT count(*) FROM messages WHERE source_id = 'item1' AND ts_utc >= ? AND ts_utc < ? "
        f"AND id IN ({involved})",
        (
            story.et("2026-03-10 00:00:00").strftime("%Y-%m-%dT%H:%M:%SZ"),
            story.et("2026-04-01 00:00:00").strftime("%Y-%m-%dT%H:%M:%SZ"),
        ),
    ).fetchone()[0]
    c03 = next(c for c in key.CLAIMS if c.claim_id == "C03")
    assert n == 13
    assert f" {n} Telegram messages" in c03.text


def test_trap_timezone_cross_device_order_flips_when_read_naively(case):
    _, conn, rendered = case
    (text_id,) = ids_for(rendered, "item1", "ns_move")
    (call_id,) = ids_for(rendered, "item2", "c_luis")
    text_raw = conn.execute("SELECT ts_raw FROM messages WHERE id = ?", (text_id,)).fetchone()[0]
    call_raw = conn.execute("SELECT ts_raw FROM calls WHERE id = ?", (call_id,)).fetchone()[0]
    text_wall, text_utc = parse_raw(text_raw)
    call_wall, call_utc = parse_raw(call_raw)
    assert text_utc < call_utc  # truth: the text came first
    assert text_wall > call_wall  # naive reading of the printed times says the opposite


def test_trap_timezone_its_done_printed_as_231_am_utc(case):
    _, conn, rendered = case
    (mid,) = ids_for(rendered, "item1", "pr03")
    raw = conn.execute("SELECT ts_raw FROM messages WHERE id = ?", (mid,)).fetchone()[0]
    assert raw == "3/5/2026 2:31:00 AM(UTC+0)"
    utc = parse_raw(raw)[1]
    local = utc + timedelta(minutes=story.eastern_offset_min(utc))
    assert (local.month, local.day, local.hour, local.minute) == (3, 4, 21, 31)


def test_trap_second_alex_numbers_differ(case):
    _, conn, _ = case
    alex = conn.execute(
        "SELECT identifier FROM contacts WHERE source_id = 'item1' AND name = 'Alex'"
    ).fetchall()
    sasha = conn.execute(
        "SELECT identifier FROM contacts WHERE source_id = 'item2' AND name = 'Sasha N'"
    ).fetchall()
    assert alex == [("+12125550182",)]
    assert set(sasha) == {("+12125550147",), ("5551234",)}


def test_trap_meeting_place_proposed_by_reyes(case):
    _, conn, rendered = case
    (mid,) = ids_for(rendered, "item1", "pr05")
    direction, body = conn.execute(
        "SELECT direction, body FROM messages WHERE id = ?", (mid,)
    ).fetchone()
    assert direction == "incoming"
    assert "kings plaza" in body


def test_trap_shared_account_has_second_author_signal(case):
    _, conn, _ = case
    bodies = [
        b
        for (b,) in conn.execute(
            "SELECT m.body FROM messages m JOIN accounts a ON a.id = m.sender_account_id "
            "WHERE m.source_id = 'item1' AND a.identifier = 'dp_garage'"
        )
    ]
    assert any("its ilya" in b for b in bodies)
    assert "got the money, come get it" in bodies


def test_filler_cannot_touch_traps(case):
    _, _, rendered = case
    trap_accounts = {
        "pet_sms",
        "pet_wa",
        "pet_tg",
        "dpg_ig",
        "rey_sms",
        "rey_wa",
        "rey_tg",
        "rey_ig",
        "ns_tg",
        "turner_sms",
    }
    for r in rendered:
        for m in r.messages:
            if not m["key"].startswith("f:"):
                continue
            body = m["body"].lower()
            assert not any(w in body for w in story.RESERVED_WORDS), (m["key"], body)
            other = m["key"].split(":")[2]
            assert other not in trap_accounts


def test_trap_timezone_order_needs_both_phones(case):
    _, conn, _ = case
    by_source = dict(
        conn.execute(
            "SELECT source_id, count(*) FROM messages WHERE body = 'move it tonight' GROUP BY 1"
        ).fetchall()
    )
    luis_calls = dict(
        conn.execute(
            "SELECT c.source_id, count(*) FROM calls c JOIN accounts a ON a.id = c.from_account_id "
            "WHERE a.identifier = '+12125550177' AND c.ts_utc LIKE '2026-03-15T02:05%' GROUP BY 1"
        ).fetchall()
    )
    assert by_source == {"item1": 1}
    assert luis_calls == {"item2": 1}


def test_trap_meeting_place_named_only_by_reyes(case):
    _, conn, _ = case
    rows = conn.execute(
        "SELECT source_id, direction FROM messages WHERE lower(body) LIKE '%kings plaza%'"
    ).fetchall()
    assert sorted(rows) == [("item1", "incoming"), ("item2", "outgoing")]


def test_trap_second_alex_hard_negative_needs_both_phones(case):
    _, conn, rendered = case
    (intro,) = ids_for(rendered, "item1", "at00")
    (asks,) = ids_for(rendered, "item2", "rn07")
    body = conn.execute("SELECT body FROM messages WHERE id = ?", (intro,)).fetchone()[0]
    sender = conn.execute(
        "SELECT a.identifier FROM messages m JOIN accounts a ON a.id = m.sender_account_id "
        "WHERE m.id = ?",
        (intro,),
    ).fetchone()[0]
    assert "alex turner" in body and sender == "+12125550182"
    asker, text = conn.execute(
        "SELECT a.identifier, m.body FROM messages m JOIN accounts a "
        "ON a.id = m.sender_account_id WHERE m.id = ?",
        (asks,),
    ).fetchone()
    assert asker == "5551234" and "who is alex turner" in text
    # neither half is on the other phone
    assert not conn.execute(
        "SELECT 1 FROM messages WHERE source_id = 'item1' AND body LIKE '%who is alex turner%'"
    ).fetchone()
    assert not conn.execute(
        "SELECT 1 FROM messages WHERE source_id = 'item2' AND body LIKE '%its alex turner%'"
    ).fetchone()
