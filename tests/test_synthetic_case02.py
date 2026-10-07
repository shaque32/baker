"""HIDDEN synthetic case02: determinism, importer agreement, draft key integrity, planted traps.

Builder threads must not read this file, eval/synthetic/case02* or eval/out/case02/.
"""

import json
import re
import sqlite3
from datetime import UTC, datetime, timedelta
from pathlib import Path

import openpyxl
import pytest

from core.contracts import GoldClaim
from core.db import apply_schema
from core.ingest.cellebrite_excel import CellebriteExcelImporter
from core.ingest.govdoc import PdfGovDocIngester
from eval.run_eval import load_jsonl
from eval.synthetic import case01_story
from eval.synthetic import case02_key as key
from eval.synthetic import case02_story as story
from eval.synthetic.case02_generate import (
    DEFAULT_OUT,
    IMPORTED_AT,
    affidavit_paragraphs,
    generate,
)

COMMITTED = Path(__file__).resolve().parents[1] / DEFAULT_OUT
COMMITTED_FILES = (
    "case.json",  # pins the SHA-256 of item1.xlsx and item2.xlsx
    "affidavit_draft.md",
    "draft_gold.jsonl",
    "ANSWER_KEY_DRAFT.md",
)
RAW_TS = re.compile(r"^(\d+)/(\d+)/(\d{4}) (\d+):(\d\d):(\d\d) (AM|PM)\(UTC([+-])(\d+)\)$")
RAW_TS_IN_TEXT = re.compile(r"\d+/\d+/\d{4} \d+:\d\d:\d\d [AP]M\(UTC[+-]\d+\)")
TABLES = (
    "sources",
    "devices",
    "accounts",
    "threads",
    "messages",
    "message_recipients",
    "attachments",
    "calls",
    "contacts",
)
HALE_WA, QUINTERO_WA = "13125550131@s.whatsapp.net", "13125550125@s.whatsapp.net"


@pytest.fixture(scope="module")
def case(tmp_path_factory):
    out = tmp_path_factory.mktemp("case02")
    db = out / "db" / "case02.db"
    rendered = generate(out, db)
    conn = sqlite3.connect(db)
    yield out, conn, rendered
    conn.close()


@pytest.fixture(scope="module")
def gold(case) -> list[GoldClaim]:
    out, _, _ = case
    return load_jsonl(out / "draft_gold.jsonl", GoldClaim)


def ids_for(rendered, device: str, story_key: str) -> list[str]:
    r = next(x for x in rendered if x.device.key == device)
    return r.keys[story_key]


def one_id(rendered, device: str, story_key: str) -> str:
    (rid,) = ids_for(rendered, device, story_key)
    return rid


def parse_raw(raw: str) -> tuple[datetime, datetime]:
    """Printed time -> (naive printed wall clock, true UTC)."""
    m = RAW_TS.match(raw)
    assert m, raw
    mo, d, y, h, mi, s, ampm, sign, off = m.groups()
    hour = int(h) % 12 + (12 if ampm == "PM" else 0)
    wall = datetime(int(y), int(mo), int(d), hour, int(mi), int(s))
    offset = int(off) * (1 if sign == "+" else -1)
    return wall, (wall - timedelta(hours=offset)).replace(tzinfo=UTC)


def central_of(raw: str) -> datetime:
    return story.central(parse_raw(raw)[1])


def utc_iso(local_central: str) -> str:
    return story.ct(local_central).strftime("%Y-%m-%dT%H:%M:%SZ")


def raw_of(conn, rid: str) -> str:
    row = conn.execute(
        "SELECT ts_raw FROM messages WHERE id = ? UNION ALL SELECT ts_raw FROM calls WHERE id = ?",
        (rid, rid),
    ).fetchone()
    return row[0]


def involving(identifier: str) -> str:
    """Message ids whose sender or a recipient has this identifier."""
    return (
        "SELECT m.id FROM messages m JOIN accounts a ON a.id = m.sender_account_id "
        f"WHERE a.identifier = '{identifier}' UNION SELECT r.message_id FROM message_recipients r "
        f"JOIN accounts a ON a.id = r.account_id WHERE a.identifier = '{identifier}'"
    )


# ---------------------------------------------------------------- determinism and size


@pytest.mark.parametrize("name", COMMITTED_FILES)
def test_committed_outputs_match_a_fresh_generation(case, name):
    out, _, _ = case
    assert (out / name).read_bytes() == (COMMITTED / name).read_bytes(), (
        f"{name} is stale: run `python -m eval.synthetic.case02_generate "
        "--db eval/out/case02/case02.db` and commit the result"
    )


def test_case_json_is_marked_hidden(case):
    out, _, _ = case
    meta = json.loads((out / "case.json").read_text(encoding="utf-8"))
    assert list(meta)[:2] == ["hidden", "note"]
    assert meta["hidden"] is True
    assert "must not read" in meta["note"]


def test_size_apps_and_languages(case):
    _, conn, rendered = case
    total = conn.execute("SELECT count(*) FROM messages").fetchone()[0]
    es = sum(m["lang"] == "es" for r in rendered for m in r.messages)
    apps = {a for (a,) in conn.execute("SELECT DISTINCT app FROM threads")}
    assert 1000 <= total <= 1500
    assert es >= 150
    assert apps == {"SMS", "WhatsApp", "Telegram", "Instagram"}
    assert conn.execute("SELECT count(DISTINCT source_id) FROM messages").fetchone()[0] == 2


def test_foreign_keys_hold(case):
    _, conn, _ = case
    assert conn.execute("PRAGMA foreign_key_check").fetchall() == []


def test_expected_db_loads_with_core_schema(case):
    _, conn, _ = case
    fresh = sqlite3.connect(":memory:")
    apply_schema(fresh)
    want = {r[0] for r in fresh.execute("SELECT name FROM sqlite_master WHERE type = 'table'")}
    have = {r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type = 'table'")}
    assert want == have


# ---------------------------------------------------------------- report <-> database


def test_the_cellebrite_importer_reproduces_the_expected_database(case):
    """Same layout as case01: the existing importer reads case02 unchanged."""
    out, expected, _ = case
    got = sqlite3.connect(":memory:")
    got.execute("PRAGMA foreign_keys = ON")
    apply_schema(got)
    for src in ("item1", "item2"):
        CellebriteExcelImporter(source_id=src, now=IMPORTED_AT).import_source(
            out / f"{src}.xlsx", got
        )
    for table in TABLES:
        a = set(expected.execute(f"SELECT * FROM {table}").fetchall())
        b = set(got.execute(f"SELECT * FROM {table}").fetchall())
        assert a == b, f"{table}: importer and expected database differ"


@pytest.mark.parametrize("source", ["item1", "item2"])
def test_database_rows_match_the_excel_report(case, source):
    out, conn, _ = case
    wb = openpyxl.load_workbook(out / f"{source}.xlsx", read_only=True)
    chats = {i: row for i, row in enumerate(wb["Chats"].iter_rows(values_only=True), start=1)}
    summary = {r[0]: r[1] for r in wb["Summary"].iter_rows(values_only=True) if len(r) > 1}
    wb.close()
    header = list(chats[2])
    body_col, ts_col = header.index("Body"), header.index("Timestamp")
    rows = conn.execute(
        "SELECT id, locator, body, ts_raw FROM messages WHERE source_id = ?", (source,)
    ).fetchall()
    assert len(rows) == len(chats) - 2
    for mid, locator, body, ts_raw in rows:
        assert mid == f"msg:{source}:{locator}"
        row = chats[int(locator.split("!")[1])]
        assert (row[body_col] or "") == body
        assert row[ts_col] == ts_raw
    assert summary["Device time zone"] == "America/Chicago"


def test_printed_times_round_trip_to_stored_utc(case):
    _, conn, _ = case
    for table in ("messages", "calls"):
        for ts_raw, ts_utc, off in conn.execute(
            f"SELECT ts_raw, ts_utc, ts_offset_min FROM {table}"
        ):
            wall, utc = parse_raw(ts_raw)
            assert utc.strftime("%Y-%m-%dT%H:%M:%SZ") == ts_utc
            assert (wall - utc.replace(tzinfo=None)) == timedelta(minutes=off)


def test_report_offsets_utc_central_dst_and_travel(case):
    _, conn, _ = case
    offsets = {
        src: {int(o) for o in offs.split(",")}
        for src, offs in conn.execute(
            "SELECT source_id, group_concat(DISTINCT ts_offset_min) FROM messages GROUP BY 1"
        )
    }
    assert offsets == {"item1": {0}, "item2": {-360, -300, -420}}
    start = story.TRAVEL_START.strftime("%Y-%m-%dT%H:%M:%SZ")
    end = story.TRAVEL_END.strftime("%Y-%m-%dT%H:%M:%SZ")
    for table in ("messages", "calls"):
        for ts_utc, off in conn.execute(
            f"SELECT ts_utc, ts_offset_min FROM {table} WHERE source_id = 'item2'"
        ):
            assert (off == -420) == (start <= ts_utc < end), (table, ts_utc, off)


# ---------------------------------------------------------------- draft answer key


def test_draft_key_shape(gold):
    assert len(gold) == 20
    assert [g.claim_id for g in gold] == [f"C{i:02d}" for i in range(1, 21)]
    split = {
        v: sum(g.gold_verdict == v for g in gold) for v in ("supported", "contradicted", "unproven")
    }
    assert split == {"supported": 7, "contradicted": 6, "unproven": 7}
    assert sorted(g.claim_id for g in gold if g.cross_device) == ["C02", "C11", "C15"]
    assert all(g.para_no >= key.FIRST_PARA for g in gold)


def test_draft_key_is_never_marked_final(case, gold):
    out, _, _ = case
    assert all(g.labeled_by == key.DRAFT_LABEL and "DRAFT" in g.labeled_by for g in gold)
    text = (out / "ANSWER_KEY_DRAFT.md").read_text(encoding="utf-8")
    assert text.startswith("# case02 answer key: DRAFT")
    assert text.count("Arsh's decision: [ ] agree") == 20
    new_rules = [r for r in key.LABELING_RULES if r.startswith("NEW (case02)")]
    assert len(new_rules) == len(key.LABELING_RULES) - 4


def test_every_cited_record_exists(case, gold):
    _, conn, _ = case
    for g in gold:
        assert g.key_evidence, g.claim_id
        for rid in g.key_evidence:
            hit = conn.execute(
                "SELECT 1 FROM messages WHERE id = ?1 UNION ALL SELECT 1 FROM calls WHERE id = ?1 "
                "UNION ALL SELECT 1 FROM contacts WHERE id = ?1 "
                "UNION ALL SELECT 1 FROM attachments WHERE id = ?1 "
                "UNION ALL SELECT 1 FROM accounts WHERE id = ?1",
                (rid,),
            ).fetchone()
            assert hit, f"{g.claim_id}: {rid} missing"


def test_claims_are_verbatim_in_their_affidavit_paragraph(case, gold):
    out, _, _ = case
    paras = {(p, n): text for p, n, text in affidavit_paragraphs()}
    md = (out / "affidavit_draft.md").read_text(encoding="utf-8")
    for g in gold:
        assert g.text in paras[(g.page, g.para_no)], g.claim_id
        assert f"{g.para_no}. {paras[(g.page, g.para_no)]}" in md
    assert "Central Time" in paras[(1, 23)]


def test_affidavit_pdf_puts_each_claim_on_its_page(case, gold):
    out, _, _ = case
    pdf = out / "db" / "affidavit.pdf"
    parsed = PdfGovDocIngester().parse(pdf)
    paras = [(p.page, " ".join(p.text.split())) for p in parsed.paragraphs]
    for g in gold:
        pages = [pg for pg, t in paras if t.startswith(f"{g.para_no}. ") and g.text in t]
        assert pages == [g.page], g.claim_id


def test_quoted_text_in_claims_matches_cited_records_verbatim(case, gold):
    _, conn, _ = case
    for g in gold:
        cited = []
        for rid in g.key_evidence:
            for sql in (
                "SELECT body FROM messages WHERE id = ?",
                "SELECT t.title FROM messages m JOIN threads t ON t.id = m.thread_id "
                "WHERE m.id = ?",
                "SELECT name FROM contacts WHERE id = ?",
                "SELECT file_name FROM attachments WHERE id = ?",
            ):
                cited += [r[0] for r in conn.execute(sql, (rid,)) if r[0]]
        for quote in re.findall(r'"([^"]+)"', g.text):
            assert any(quote in c for c in cited), f"{g.claim_id}: {quote!r} not in cited records"


def test_printed_times_in_the_key_are_real_cited_values(case, gold):
    _, conn, _ = case
    for g in gold:
        cited = {raw_of(conn, rid) for rid in g.key_evidence if rid.startswith(("msg:", "call:"))}
        for printed in RAW_TS_IN_TEXT.findall(g.rationale):
            assert printed in cited, f"{g.claim_id}: {printed} is not a cited record's time"


# ---------------------------------------------------------------- planted traps


def test_trap_travel_tz_message_dates_differ_by_zone(case):
    _, conn, rendered = case
    qh03 = raw_of(conn, one_id(rendered, "item2", "qh03"))
    assert qh03 == "3/19/2026 10:30:15 PM(UTC-7)"
    assert central_of(qh03) == datetime(2026, 3, 20, 0, 30, 15)  # C17 holds in Central
    assert "qh03" not in rendered[0].keys  # Item 2 only: the per-row offset is the only route
    s07_item2 = raw_of(conn, one_id(rendered, "item2", "s07"))
    s07_item1 = raw_of(conn, one_id(rendered, "item1", "s07"))
    assert s07_item2 == "3/20/2026 11:41:10 PM(UTC-7)"  # before midnight where the phone was
    assert s07_item1 == "3/21/2026 6:41:10 AM(UTC+0)"
    assert central_of(s07_item2) == datetime(2026, 3, 21, 1, 41, 10)  # C18 fails in Central


def test_trap_dst_needs_the_summer_offset(case):
    _, conn, rendered = case
    raw = raw_of(conn, one_id(rendered, "item1", "s03"))
    assert raw == "3/10/2026 1:15:33 AM(UTC+0)"
    assert central_of(raw) == datetime(2026, 3, 9, 20, 15, 33)
    assert story.central_offset_min(parse_raw(raw)[1]) == -300  # CST would give 7:15 PM


def test_trap_group_sender_is_hale_not_brandt(case):
    _, conn, rendered = case
    mid = one_id(rendered, "item1", "g05")
    sender, direction = conn.execute(
        "SELECT a.identifier, m.direction FROM messages m JOIN accounts a "
        "ON a.id = m.sender_account_id WHERE m.id = ?",
        (mid,),
    ).fetchone()
    assert sender == HALE_WA and direction == "incoming"
    body = conn.execute("SELECT body FROM messages WHERE id = ?", (mid,)).fetchone()[0]
    assert conn.execute(
        "SELECT count(*) FROM messages WHERE body = ? AND direction = 'outgoing'", (body,)
    ).fetchone() == (0,)
    recips = conn.execute(
        "SELECT count(*) FROM message_recipients WHERE message_id = ?", (mid,)
    ).fetchone()[0]
    assert recips == 3  # a group of four


def test_trap_group_is_left_out_of_the_item2_report(case):
    _, conn, _ = case
    titles = dict(
        conn.execute(
            "SELECT source_id, count(*) FROM threads WHERE title = 'Westside Flips' GROUP BY 1"
        ).fetchall()
    )
    assert titles == {"item1": 1}
    assert conn.execute(
        "SELECT count(*) FROM accounts WHERE source_id = 'item2' AND device_id IS NOT NULL "
        "AND identifier = ?",
        (QUINTERO_WA,),
    ).fetchone() == (1,)


def test_trap_quoted_speech_attributes_the_words_to_wade(case):
    _, conn, rendered = case
    body, direction = conn.execute(
        "SELECT body, direction FROM messages WHERE id = ?", (one_id(rendered, "item1", "s05"),)
    ).fetchone()
    assert body == 'wade said "bring the cash friday"' and direction == "outgoing"


def test_trap_negation_only_brandt_text_to_mercer_that_day_is_a_denial(case):
    _, conn, rendered = case
    rows = conn.execute(
        "SELECT m.id, m.body FROM messages m WHERE m.source_id = 'item1' "
        "AND m.direction = 'outgoing' AND m.ts_utc >= ? AND m.ts_utc < ? "
        f"AND m.id IN ({involving('+13125550150')})",
        (utc_iso("2026-03-24 00:00:00"), utc_iso("2026-03-25 00:00:00")),
    ).fetchall()
    assert rows == [(one_id(rendered, "item1", "km02"), "no. i never sold chino anything. drop it")]


def test_trap_call_count_five_records_two_connected(case):
    _, conn, rendered = case
    calls = conn.execute(
        "SELECT c.id, c.direction, c.duration_s, c.ts_raw FROM calls c "
        "JOIN accounts f ON f.id = c.from_account_id JOIN accounts t ON t.id = c.to_account_id "
        "WHERE c.source_id = 'item1' AND '+13125550125' IN (f.identifier, t.identifier) "
        "AND c.ts_utc >= ? AND c.ts_utc < ? ORDER BY c.ts_utc",
        (utc_iso("2026-03-10 00:00:00"), utc_iso("2026-03-15 00:00:00")),
    ).fetchall()
    assert [c[0] for c in calls] == [one_id(rendered, "item1", f"k{i}") for i in range(1, 6)]
    assert [(c[1], c[2]) for c in calls] == [
        ("outgoing", 271),
        ("missed", 0),
        ("outgoing", 62),
        ("outgoing", 0),
        ("missed", 0),
    ]
    assert sum(c[2] > 0 for c in calls) == 2
    printed_days = [parse_raw(c[3])[0].day for c in calls]
    assert printed_days == [10, 11, 13, 13, 15]  # two print outside the window's dates in UTC
    assert raw_of(conn, one_id(rendered, "item1", "k3")) == "3/13/2026 2:14:05 AM(UTC+0)"


def test_trap_deleted_flag_six_telegram_rows_with_brandt(case):
    _, conn, rendered = case
    thread = conn.execute(
        "SELECT m.id, m.deleted_flag FROM messages m JOIN threads t ON t.id = m.thread_id "
        "WHERE m.source_id = 'item2' AND t.app = 'Telegram' "
        f"AND m.id IN ({involving('8100219')})"
    ).fetchall()
    deleted = {mid for mid, flag in thread if flag == 1}
    assert len(thread) == 14
    assert deleted == {one_id(rendered, "item2", k) for k in story.ITEM2_DELETED_TG}
    all_item2 = conn.execute(
        "SELECT count(*) FROM messages WHERE source_id = 'item2' AND deleted_flag = 1"
    ).fetchone()[0]
    assert all_item2 == 9  # a naive count of every Deleted row is wrong
    item1_copies = [one_id(rendered, "item1", k) for k in story.ITEM2_DELETED_TG]
    flags = conn.execute(
        f"SELECT DISTINCT deleted_flag FROM messages WHERE id IN ({','.join('?' * 6)})",
        item1_copies,
    ).fetchall()
    assert flags == [(0,)]
    c10 = next(c for c in key.CLAIMS if c.claim_id == "C10")
    assert "six Telegram messages" in c10.text


def test_trap_absence_contradicted_only_on_the_other_phone(case):
    _, conn, rendered = case
    since = utc_iso("2026-03-22 00:00:00")
    item2_msgs = conn.execute(
        "SELECT count(*) FROM messages WHERE source_id = 'item2' AND ts_utc >= ? "
        f"AND id IN ({involving(HALE_WA)})",
        (since,),
    ).fetchone()[0]
    item2_calls = conn.execute(
        "SELECT count(*) FROM calls c JOIN accounts f ON f.id = c.from_account_id "
        "JOIN accounts t ON t.id = c.to_account_id WHERE c.source_id = 'item2' "
        "AND c.ts_utc >= ? AND '+13125550131' IN (f.identifier, t.identifier)",
        (since,),
    ).fetchone()[0]
    assert (item2_msgs, item2_calls) == (0, 0)  # Item 2 alone looks like silence
    g07, g08 = one_id(rendered, "item1", "g07"), one_id(rendered, "item1", "g08")
    pairs = conn.execute(
        "SELECT a.identifier, m.ts_utc FROM messages m JOIN accounts a "
        "ON a.id = m.sender_account_id WHERE m.id IN (?, ?) ORDER BY m.ts_utc",
        (g07, g08),
    ).fetchall()
    assert [p[0] for p in pairs] == [QUINTERO_WA, HALE_WA]
    assert all(ts >= since for _, ts in pairs)
    hale_got_g07 = conn.execute(
        "SELECT 1 FROM message_recipients r JOIN accounts a ON a.id = r.account_id "
        "WHERE r.message_id = ? AND a.identifier = ?",
        (g07, HALE_WA),
    ).fetchone()
    assert hale_got_g07


def test_trap_absence_unproven_no_row_with_chino_0143_after_march_16(case):
    _, conn, rendered = case
    since = utc_iso("2026-03-17 00:00:00")
    msgs = conn.execute(
        f"SELECT count(*) FROM messages WHERE ts_utc >= ? AND id IN ("
        f"{involving('13125550143@s.whatsapp.net')} UNION {involving('+13125550143')})",
        (since,),
    ).fetchone()[0]
    calls = conn.execute(
        "SELECT count(*) FROM calls c JOIN accounts f ON f.id = c.from_account_id "
        "JOIN accounts t ON t.id = c.to_account_id "
        "WHERE c.ts_utc >= ? AND '+13125550143' IN (f.identifier, t.identifier)",
        (since,),
    ).fetchone()[0]
    assert (msgs, calls) == (0, 0)
    last = raw_of(conn, one_id(rendered, "item1", "ch07"))
    assert central_of(last) == datetime(2026, 3, 14, 10, 12, 19)
    fidelity = {f for (f,) in conn.execute("SELECT fidelity FROM sources")}
    assert fidelity == {"curated_report"}


def test_trap_sarcasm_context(case):
    _, conn, rendered = case
    bodies = [
        conn.execute(
            "SELECT body FROM messages WHERE id = ?", (one_id(rendered, "item1", k),)
        ).fetchone()[0]
        for k in ("g01", "g02", "g03")
    ]
    assert bodies[1] == "yeah im the kingpin lol"
    assert bodies[0].endswith("lol") and bodies[2] == "jajaja"
    assert central_of(raw_of(conn, one_id(rendered, "item1", "g02"))).day == 6


def test_trap_name_collision_two_chinos_two_numbers(case):
    _, conn, _ = case
    rows = conn.execute(
        "SELECT source_id, identifier FROM contacts WHERE name = 'Chino'"
    ).fetchall()
    assert sorted(rows) == [("item1", "+13125550143"), ("item2", "+13125550158")]
    assert not conn.execute(
        "SELECT 1 FROM contacts WHERE source_id = 'item2' AND identifier = '+13125550143'"
    ).fetchone()


def test_trap_attachment_only_has_a_file_name_and_nothing_else(case):
    _, conn, rendered = case
    mid = one_id(rendered, "item1", "ch03")
    body = conn.execute("SELECT body FROM messages WHERE id = ?", (mid,)).fetchone()[0]
    att = conn.execute(
        "SELECT file_name, mime_type, sha256 FROM attachments WHERE message_id = ?", (mid,)
    ).fetchall()
    assert body == ""
    assert att == [("IMG_4471.jpg", None, None)]


def test_trap_translation_messages(case):
    _, conn, rendered = case
    r2 = next(r for r in rendered if r.device.key == "item2")
    by_key = {m["key"]: m for m in r2.messages}
    assert by_key["mq01"]["lang"] == "es" and "el miércoles me voy" in by_key["mq01"]["body"]
    assert by_key["mf01"]["lang"] == "es" and "caliente" in by_key["mf01"]["body"]
    for word in ("polic", "storage", "unit", "almacén", "bodega"):
        assert word not in by_key["mf01"]["body"].lower()
    assert datetime(2026, 3, 18).weekday() == 2  # 'el miércoles' after Sunday Mar 15 is Mar 18


def test_trap_handle_reuse_one_name_two_user_ids(case):
    out, conn, _ = case
    wb = openpyxl.load_workbook(out / "item1.xlsx", read_only=True)
    header, *rows = list(wb["Chats"].iter_rows(min_row=2, values_only=True))
    wb.close()
    col = {name: i for i, name in enumerate(header)}
    froms = {
        r[col["From"]]
        for r in rows
        if r[col["Name"]] == "Vic Tools" and r[col["Direction"]] == "Incoming"
    }
    assert froms == {"8200417 Vic Tools", "8200952 Vic Tools"}
    threads = conn.execute(
        "SELECT count(*) FROM threads WHERE source_id = 'item1' AND title = 'Vic Tools'"
    ).fetchone()[0]
    assert threads == 2


def test_filler_cannot_touch_traps(case):
    _, _, rendered = case
    trap_accounts = {
        "nb_sms",
        "nb_tg",
        "rq_sms",
        "rq_tg",
        "hale_wa",
        "salas_wa",
        "salas_sms",
        "vt_old",
        "vt_new",
    }
    for r in rendered:
        for m in r.messages:
            if not m["key"].startswith("f:"):
                continue
            body = m["body"].lower()
            assert not any(w in body for w in story.RESERVED_WORDS), (m["key"], body)
            other = m["key"].split(":")[2]
            assert other not in trap_accounts
            if m["group"]:
                assert m["ts_utc"] < "2026-03-18T05:00:00Z"
            if other == "kyle_sms":
                at = datetime.fromisoformat(m["ts_utc"].replace("Z", "+00:00"))
                assert story.central(at).date() != datetime(2026, 3, 24).date()


def test_no_case01_person_or_number_appears(case):
    out, _, _ = case
    texts = [
        (out / name).read_text(encoding="utf-8")
        for name in ("case.json", "affidavit_draft.md", "draft_gold.jsonl", "ANSWER_KEY_DRAFT.md")
    ]
    for src in ("item1", "item2"):
        wb = openpyxl.load_workbook(out / f"{src}.xlsx", read_only=True)
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                texts.append(" ".join(str(c) for c in row if c is not None))
        wb.close()
    blob = "\n".join(texts)
    names = set(case01_story.PEOPLE.values()) - {"Petrov's mother", "Reyes's mother"}
    surnames = {n.split()[-1] for n in names if " " in n and not n.startswith("unknown")}
    for n in names | surnames | {"Petrov", "Reyes", "Marc Garage", "Sasha N", "northstar"}:
        assert n not in blob, n
    for a in case01_story.ACCOUNTS.values():
        assert a.identifier not in blob, a.identifier
