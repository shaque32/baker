"""Cellebrite Physical Analyzer / Reader Excel report (.xlsx) importer.

Layout assumptions: docs/ingest/cellebrite_report.md.
"""

from __future__ import annotations

import io
import re
import sqlite3
import zipfile
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from core.contracts import Source, SourceKind
from core.ingest.cellebrite_common import (
    HEADER_SCAN_ROWS,
    NOT_IMPORTED_CATEGORIES,
    SUMMARY_SHEETS,
    Cell,
    Coverage,
    Merge,
    RawTable,
    ReportInput,
    app_for_title,
    header_preview,
    kind_for_title,
    known_header,
    norm,
    text,
    write_report,
)
from core.ingest.source_file import read_source, require_sha256

_COUNT_RE = re.compile(r"\s*\(\s*[\d,]+\s*\)$")
# Excel stores characters XML cannot hold as text, such as the carriage return of a CRLF line
# break, as _xHHHH_ (ECMA-376 Part 1, 22.9.2.19 ST_Xstring), and a literal "_x" as _x005F_x.
# openpyxl returns cell strings still escaped.
_XSTRING_RE = re.compile(r"_x([0-9A-Fa-f]{4})_")
_MERGE_RE = re.compile(rb'<mergeCell\s+ref="([A-Z]+[0-9]+:[A-Z]+[0-9]+)"')


class CellebriteExcelImporter:
    """Implements core.contracts.EvidenceImporter."""

    name = "cellebrite_excel"

    def __init__(
        self,
        source_id: str | None = None,
        now: datetime | None = None,
        date_order: str | None = None,
        expected_sha256: str | None = None,
    ) -> None:
        """`date_order` ('mdy' or 'dmy'): the report's day/month order, when a person states
        it from the report's date format. Without it, dates like 3/4/2026 are read only if
        another value in the report decides the order; otherwise they are reported unread."""
        self.source_id = source_id
        self.now = now
        self.date_order = date_order
        self.expected_sha256 = expected_sha256  # refuse other bytes (source_file.require_sha256)

    def import_source(self, path: Path, conn: sqlite3.Connection) -> Source:
        sf = read_source(path)
        require_sha256(path.name, self.expected_sha256, sf.sha256)
        report = read_workbook(sf.data)
        return write_report(
            conn,
            sf,
            SourceKind.CELLEBRITE_EXCEL,
            self.name,
            self.source_id,
            self.now or datetime.now(UTC),
            report,
            self.date_order,
        )


def _merged_ranges(data: bytes, sheet_path: str | None) -> list[tuple[int, int, int, int]]:
    """Merged ranges of one sheet as (min_col, min_row, max_col, max_row), 1-based. Read-only
    workbooks do not expose them, so they are read from the sheet's XML, where they follow the
    cell data."""
    if not sheet_path:
        return []
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            xml = z.read(sheet_path)
    except (KeyError, zipfile.BadZipFile):
        return []
    tail = xml[xml.rfind(b"</sheetData>") :]
    return [range_boundaries(m.decode()) for m in _MERGE_RE.findall(tail)]  # type: ignore[misc]


def _join_two_row_header(rows: list[tuple[str, list[Cell]]], merges: list[Merge]) -> list[str]:
    """A header written on two rows, with a group name merged across the columns it covers
    ('Timestamp' over 'Date' and 'Time'). When 'Group: Child' is a known column name, the two
    rows become one: the child row takes the joined names plus every name the group row holds
    alone, and the group row is blanked so it is not taken for the header. Returns notes."""
    notes: list[str] = []
    by_row: dict[int, list[Merge]] = {}
    for mg in merges:
        if mg.last_row == mg.first_row and mg.last_col > mg.first_col:
            by_row.setdefault(mg.first_row, []).append(mg)
    for r, row_merges in sorted(by_row.items()):
        if r >= HEADER_SCAN_ROWS or r + 1 >= len(rows):
            continue
        parent_cells, child_cells = rows[r][1], rows[r + 1][1]
        joined: dict[int, str] = {}
        groups: list[str] = []
        for mg in row_merges:
            parent = parent_cells[mg.first_col]
            if not isinstance(parent, str):
                continue
            for c in range(mg.first_col, mg.last_col + 1):
                child = child_cells[c]
                name = f"{parent.strip()}: {str(child).strip()}"
                if isinstance(child, str) and known_header(name):
                    joined[c] = name
            groups.append(parent.strip())
        if not joined:
            continue  # a title over the header row, or a merged data cell
        grouped_cols = {c for mg in row_merges for c in range(mg.first_col, mg.last_col + 1)}
        for c, name in joined.items():
            child_cells[c] = name
        for c, v in enumerate(parent_cells):
            if c not in grouped_cols and v is not None and child_cells[c] is None:
                child_cells[c] = v
        rows[r] = (rows[r][0], [None] * len(parent_cells))
        notes.append(
            f"the header on rows {rows[r][0]} and {rows[r + 1][0]} was read as one row "
            f"(grouped columns: {', '.join(repr(g) for g in groups)})."
        )
    return notes


def _fill_header_band(rows: list[tuple[str, list[Cell]]], merges: list[Merge]) -> None:
    """A header name merged down two rows sits in the top cell only; copy it into the cells
    below so the header row holds every name. Header rows only (the first rows of a sheet)."""
    for mg in merges:
        if mg.first_row >= HEADER_SCAN_ROWS or mg.last_row == mg.first_row:
            continue
        top = rows[mg.first_row][1]
        value = top[mg.first_col] if mg.first_col < len(top) else None
        if not isinstance(value, str) or not known_header(value):
            continue
        for r in range(mg.first_row + 1, min(mg.last_row, HEADER_SCAN_ROWS - 1) + 1):
            cells = rows[r][1]
            if mg.first_col < len(cells) and cells[mg.first_col] is None:
                cells[mg.first_col] = value


def _unescape(value: Cell, cov: Coverage) -> Cell:
    """A cell string as Excel means it: each _xHHHH_ escape becomes its character, in one pass,
    so _x005F_x000D_ stays the literal text '_x000D_'. Counted in the coverage record."""
    if not isinstance(value, str) or "_x" not in value:
        return value
    out = _XSTRING_RE.sub(lambda m: chr(int(m.group(1), 16)), value)
    if out != value:
        cov.counters["cells_with_excel_escapes_decoded"] += 1
    return out


def read_workbook(data: bytes) -> ReportInput:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    cov = Coverage()
    summary: list[tuple[str, str, str]] = []
    tables: list[RawTable] = []
    found_summary = False
    try:
        for ws in wb.worksheets:
            title = ws.title
            rows: list[tuple[str, list[Cell]]] = [
                (f"{title}!{i}", [_unescape(c, cov) for c in cells])
                for i, cells in enumerate(ws.iter_rows(values_only=True), start=1)
            ]
            if norm(title) in SUMMARY_SHEETS:
                found_summary = True
                for loc, cells in rows:
                    vals = [text(c) for c in cells]
                    present = [v for v in vals if v is not None]
                    if len(present) >= 2:
                        summary.append((present[0], present[1], loc))
                continue
            kind = kind_for_title(title)
            if kind is None:
                # Never guess a sheet's kind from its columns: a Timeline sheet holds the same
                # messages and calls again and would import them twice.
                probe = RawTable(name=title, kind_hint=None, rows=rows)
                cov.not_imported.append(
                    {
                        "name": title,
                        "reason": (
                            "category not imported"
                            if _COUNT_RE.sub("", norm(title)) in NOT_IMPORTED_CATEGORIES
                            else "sheet not recognised"
                        ),
                        "rows": sum(1 for _l, c in rows if any(text(x) is not None for x in c)),
                        "columns": header_preview(probe),
                    }
                )
                continue
            if getattr(ws, "sheet_state", "visible") != "visible":
                cov.notes.append(f"Sheet {title!r} is hidden in the workbook; it was read.")
                cov.hidden_sheets.append(title)
            merges = [
                Merge(min_row - 1, max_row - 1, min_col - 1, max_col - 1)
                for (min_col, min_row, max_col, max_row) in _merged_ranges(
                    data, getattr(ws, "_worksheet_path", None)
                )
            ]
            width = max((len(c) for _l, c in rows), default=0)
            rows = [(loc, cells + [None] * (width - len(cells))) for loc, cells in rows]
            _fill_header_band(rows, merges)
            notes = _join_two_row_header(rows, merges)
            tables.append(
                RawTable(
                    name=title,
                    kind_hint=kind,
                    rows=rows,
                    app_hint=app_for_title(title),
                    merges=merges,
                    header_notes=notes,
                )
            )
    finally:
        wb.close()
    cov.summary_found = found_summary
    return ReportInput(summary=summary, tables=tables, coverage=cov)
