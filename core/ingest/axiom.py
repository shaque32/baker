"""Magnet AXIOM report export importer (XML and Excel). UNVERIFIED format support.

An examiner using Magnet AXIOM Examine hands over a report export. Baker reads two of them:

- the XML report: `Artifact` elements (one per artifact type, named like 'Android WhatsApp
  Messages'), each holding `Hits/Hit` records made of named `Fragment` values;
- the Excel report (.xlsx): one sheet per artifact type, one header row of the same fragment
  names, one row per hit.

Both were built from public descriptions only: Magnet's documentation of its export formats and
an open-source parser of AXIOM XML reports (casework/CASE-Implementation-AXIOM), which names the
artifact types and fragments. Baker has never been given a real AXIOM export, so every import says
it was read with an unverified reader, on the case screen, in `baker import` and in the report's
limitations. Layout assumptions and the steps to verify: docs/ingest/axiom.md.

What it shares with the Cellebrite importers (core/ingest/cellebrite_common.py): the record
mapping for messages and calls (threads, senders, recipients, direction, owner accounts), time
parsing (no guessed day/month order, no guessed zone), the stable id scheme and the coverage
record. What differs: the artifact name gives the record kind and the app, a sender or partner
comes with its name in a separate field, and a time field's name states its UTC offset and date
format ('Message Sent Date/Time - UTC+00:00 (M/d/yyyy)').

A report export holds what the examiner chose to export, so it imports as a curated report: it
can never support a claim that something is absent from the phone.

The input is hashed and read once; an XML report is parsed in chunks (only the records are kept,
never the whole document) and hashed again before anything is stored.
"""

from __future__ import annotations

import hashlib
import io
import json
import re
import sqlite3
from collections import Counter
from collections.abc import Callable
from dataclasses import dataclass, field, replace
from datetime import UTC, datetime
from pathlib import Path
from typing import IO
from xml.parsers import expat

from openpyxl import load_workbook

from core.contracts import Fidelity, Source, SourceKind
from core.encryption import open_input
from core.ingest import forwarded
from core.ingest.cellebrite_common import (
    HEADER_SCAN_ROWS,
    Cell,
    Coverage,
    DateOrder,
    Party,
    RecordKind,
    Row,
    TableCoverage,
    TimeContext,
    _accounts,
    _calls,
    _contacts,
    _listed,
    _messages,
    _Writer,
    coverage_payload,
    extraction_type,
    norm,
    parse_parties,
    parse_time,
    text,
)
from core.ingest.source_file import require_sha256
from core.review import audit_log

IMPORTER_NAME = "axiom"
IMPORTER_VERSION = "0.2.0"  # 0.2.0: forwarded and quoted messages marked, copies imported
# Becomes "verified" only after a real AXIOM export has been checked against AXIOM Examine
# (docs/ingest/axiom.md, "Before calling it verified"). Every import records it.
FORMAT_SUPPORT = "unverified"
CHUNK = 1 << 20

FIDELITY_REASON = (
    "Magnet AXIOM report export: the examiner chooses which artifacts and records it holds, so "
    "it cannot support a claim that something is absent from the device."
)

# ---------------------------------------------------------------- artifact names

# Words around the app in an artifact name: a platform in front ('Android', 'iOS'), a platform
# after a dash ('Signal Messages - iOS'), the record kind at the end ('WhatsApp Messages').
_PLATFORM_RE = re.compile(r"^(?:android|ios|apple|windows|macos|mac)\s+", re.I)
_PLATFORM_TAIL_RE = re.compile(r"\s+-\s+(?:android|ios|windows|macos|mac)$", re.I)
_COUNT_RE = re.compile(r"\s*\(\s*[\d,]+\s*\)$")  # 'Android SMS (1,204)'
# (kind, endings of the lowercased name). The first ending that matches decides.
_KIND_ENDINGS: list[tuple[str, tuple[str, ...]]] = [
    ("device", ("device information",)),
    ("accounts", ("accounts information", "account information", "user accounts", "accounts")),
    ("calls", ("call logs", "call log", "call history", "calls")),
    ("contacts", ("contacts",)),
    ("chats", ("chat messages", "direct messages", "messages", "chats", "chat")),
]
_SMS_WORDS = ("sms", "mms", "imessage")
# Artifact types Baker knows and deliberately does not import (yet), matched on words in the
# name. Listed in coverage as categories not imported; any other name is listed as not recognised.
NOT_IMPORTED_WORDS = (
    "email", "mail", "web", "history", "bookmark", "cookie", "search", "location", "locations",
    "picture", "pictures", "video", "videos", "audio", "documents", "calendar", "note", "notes",
    "wifi", "wi-fi", "bluetooth", "cell tower", "timeline", "application usage", "knowledgec",
    "installed", "voicemail", "recordings", "photos", "media", "file system", "sim card",
    "password", "credentials", "notifications", "screen time", "health", "journey",
)  # fmt: skip
SUMMARY_SHEETS = {"case information", "case summary", "summary", "report summary"}


EXCEL_SHEET_NAME_LIMIT = 31  # characters; Excel cuts longer sheet names


def classify(name: str, truncated: bool = False) -> tuple[str | None, str | None]:
    """(kind, app) for an artifact name. kind is a RecordKind value, 'device' (the phone's own
    details), or None when Baker does not import the artifact. SMS-family artifacts are 'sms'
    with the app 'SMS'; a call log without an app is the phone's ('Phone').

    With `truncated` (an Excel sheet name at the 31-character limit), the name may end partway
    through its kind words ('Android WhatsApp Accounts Infor'); four or more letters of the
    ending are then enough."""
    full = _COUNT_RE.sub("", norm(name))
    n = _PLATFORM_RE.sub("", _PLATFORM_TAIL_RE.sub("", full).strip() + " ").strip()
    for kind, endings in _KIND_ENDINGS:
        for end in endings:
            cuts = range(len(end), 3, -1) if truncated else (len(end),)
            for k in cuts:
                part = end[:k]
                if n == part or n.endswith(" " + part):
                    rest = n[: len(n) - len(part)].strip()
                    words = set(re.split(r"[\s/]+", rest))
                    if kind == "chats" and (words & set(_SMS_WORDS) or rest in ("", "text")):
                        return RecordKind.SMS.value, "SMS"
                    if kind == "calls" and not rest:
                        return kind, "Phone"
                    return kind, _app_name(name, rest) if rest else None
    words = set(re.split(r"[\s/]+", n))
    if words & set(_SMS_WORDS) or "sms/mms" in n:
        return RecordKind.SMS.value, "SMS"
    return None, None


def _app_name(original: str, rest_lower: str) -> str:
    """The app as the artifact name spells it ('Android WhatsApp Messages' -> 'WhatsApp')."""
    s = _COUNT_RE.sub("", original).strip()
    s = _PLATFORM_TAIL_RE.sub("", s).strip()
    s = _PLATFORM_RE.sub("", s).strip()
    words = s.split()
    return " ".join(words[: len(rest_lower.split())])


def not_imported_reason(name: str) -> str:
    n = norm(name)
    if any(re.search(rf"(?:^|[\s/]){re.escape(w)}s?(?:$|[\s/])", n) for w in NOT_IMPORTED_WORDS):
        return "category not imported"
    return "artifact not recognised"


# ---------------------------------------------------------------- field names

# A time field's name carries its zone and date format: 'Message Date/Time - UTC-05:00 (M/d/yyyy)'.
_TIME_NAME_RE = re.compile(r"^(?P<base>.+?)\s+-\s+(?P<zone>[^()]+?)\s*(?:\((?P<fmt>[^()]*)\))?\s*$")
_UTC_ZONE_RE = re.compile(r"^(?:UTC|GMT)\s*(?P<off>[+-]\s*\d{1,2}(?::?\d{2})?)?$", re.I)


@dataclass(frozen=True)
class FieldName:
    written: str  # as the report writes it
    base: str  # lowercased, the zone and format removed
    offset: int | None = None  # minutes, when the name states a UTC offset
    zone: str | None = None  # a zone the name states that is not a UTC offset (not applied)
    date_format: str | None = None  # 'M/d/yyyy' as written


def field_name(written: str) -> FieldName:
    w = " ".join(str(written).split())
    m = _TIME_NAME_RE.match(w)
    if m and re.search(r"date|time", m.group("base"), re.I):
        zone = m.group("zone").strip()
        z = _UTC_ZONE_RE.match(zone)
        off = None
        if z:
            raw = (z.group("off") or "+00:00").replace(" ", "")
            sign = -1 if raw[0] == "-" else 1
            digits = raw[1:].replace(":", "")
            h, mi = (int(digits), 0) if len(digits) <= 2 else (int(digits[:-2]), int(digits[-2:]))
            off = sign * (h * 60 + mi)
        fmt = (m.group("fmt") or "").strip() or None
        return FieldName(written, norm(m.group("base")), off, None if z else zone, fmt)
    return FieldName(written, norm(w))


def date_format_order(fmt: str | None) -> str | None:
    """'M/d/yyyy' -> 'mdy', 'd/M/yyyy' and 'dd.MM.yyyy' -> 'dmy', 'yyyy-MM-dd' -> 'ymd'."""
    if not fmt:
        return None
    letters = [c for c in fmt if c in "yMd"]
    order = "".join(dict.fromkeys(letters))
    return {"Mdy": "mdy", "dMy": "dmy", "yMd": "ymd"}.get(order)


# Field aliases per record kind, lowercased; the first field present wins. Times are listed in
# order of preference: a message's sent time comes before its received time.
_CHAT_FIELDS: dict[str, tuple[str, ...]] = {
    "from_id": ("sender", "sender id", "from", "author"),
    "from_name": ("sender name", "from name", "author name"),
    "to_id": ("recipient(s)", "recipients", "recipient", "receiver", "to"),
    "to_name": ("recipient name", "recipient name(s)", "receiver name", "to name"),
    "partner_id": ("partner", "partners", "participant information", "participants"),
    "partner_name": ("partner name", "participant name(s)", "participant names"),
    "chat_id": ("chat id", "conversation id", "thread id", "chat identifier"),
    "chat_name": ("conversation name", "chat name", "group name", "conversation"),
    "body": ("message", "message body", "body", "text"),
    "ts": (
        "message date/time", "message sent date/time", "sent date/time",
        "message received date/time", "received date/time", "created date/time",
        "date/time", "timestamp date/time",
    ),
    "direction": ("direction", "message direction"),
    "status": ("message status", "status"),
    "deleted": ("deleted", "deleted state", "record status"),
    "tag": ("tags", "tag"),
    "attachment": ("attachment", "attachments", "attachment(s)", "attachment name",
                   "attachment file name", "attachment name(s)"),
    "recovery": ("recovery method",),
    # Forwarded and quoted messages (pilot gate 4). ASSUMED names, conservative: no real AXIOM
    # export has been checked (docs/ingest/axiom.md). AXIOM's chat artifacts are described as
    # carrying a forwarded flag and, for a reply, the quoted message's text and sender; Baker
    # reads only these exact names (lowercased):
    # - a flag that the message was forwarded ('Yes', 'True', '1', 'Forwarded'; 'No', 'False',
    #   '0' and 'Not forwarded' mean not; any other value fails closed and marks it forwarded);
    "forwarded": ("forwarded", "is forwarded"),
    # - the forwarded message's original author, stored as written, never resolved;
    "forwarded_from": ("forwarded from", "original sender", "original author"),
    # - the quoted (replied-to) or forwarded original carried inside the message, imported one
    #   level deep as its own record at '<locator>/Embedded[1]', with its own sender as written.
    "quoted_body": ("quoted message", "quoted message body", "quoted text"),
    "quoted_from_id": ("quoted message sender", "quoted message author", "quoted sender"),
    "quoted_from_name": ("quoted message sender name", "quoted sender name"),
}  # fmt: skip
_SMS_FIELDS = {**_CHAT_FIELDS, "type": ("type", "message type")}
_CALL_FIELDS: dict[str, tuple[str, ...]] = {
    "partner_id": ("partner", "partners", "phone number", "number", "partner(s)"),
    "partner_name": ("partner name", "contact name", "name"),
    "from_id": ("caller", "from"),
    "to_id": ("callee", "to"),
    "ts": ("call date/time", "start date/time", "date/time", "call start date/time"),
    "duration": ("call duration", "duration"),
    "direction": ("direction", "call direction", "call type", "type"),
    "status": ("call status", "status"),
    "deleted": ("deleted", "deleted state", "record status"),
    "recovery": ("recovery method",),
}
_CONTACT_FIELDS: dict[str, tuple[str, ...]] = {
    "name": ("display name", "name", "full name", "contact name"),
    "first_name": ("first name",),
    "last_name": ("last name",),
    "phone": ("phone number(s)", "phone numbers", "phone number", "phone"),
    "email": ("email address(es)", "email addresses", "email address", "email(s)", "email"),
    "username": ("user name", "username", "user id", "account"),
    "deleted": ("deleted", "deleted state", "record status"),
    "recovery": ("recovery method",),
}
_ACCOUNT_FIELDS: dict[str, tuple[str, ...]] = {
    "username": ("phone number", "user name", "username", "user id", "account id", "account",
                 "email address", "email"),
    "name": ("whatsapp name", "display name", "full name", "name", "account name"),
    "recovery": ("recovery method",),
}  # fmt: skip
_DEVICE_FIELDS: dict[str, tuple[str, ...]] = {
    "device": ("device name", "model id", "model", "device model"),
    "manufacturer": ("manufacturer", "make"),
    "os": ("os version", "operating system", "os"),
    "time_zone": ("time zone", "timezone", "device time zone"),
}
FIELDS = {
    RecordKind.CHATS.value: _CHAT_FIELDS,
    RecordKind.SMS.value: _SMS_FIELDS,
    RecordKind.CALLS.value: _CALL_FIELDS,
    RecordKind.CONTACTS.value: _CONTACT_FIELDS,
    RecordKind.ACCOUNTS.value: _ACCOUNT_FIELDS,
    "device": _DEVICE_FIELDS,
}
# Provenance fields every artifact carries. Known, and not read into the record: 'Source' is the
# file inside the image the hit came from (never the app), 'Location' the place inside that file.
IGNORED_FIELDS = {"source", "location", "evidence number", "item id", "comments", "artifact type"}
# Values of 'Recovery method' that mean the record was rebuilt from raw or unallocated data.
_CARVED_RE = re.compile(r"carv", re.I)
# Summary keys of an Excel report's case sheet.
SUMMARY_KEYS: dict[str, tuple[str, ...]] = {
    "extraction_start": (
        "acquisition date/time", "acquisition date", "extraction date/time",
        "extraction start date/time", "extraction date", "image acquired",
    ),
    "extraction_type": ("acquisition type", "extraction type", "image type"),
    "device": ("device", "device name", "device model", "model"),
    "os": ("os version", "operating system"),
    "tool_version": ("axiom version", "magnet axiom version", "examine version", "version"),
}  # fmt: skip


# ---------------------------------------------------------------- reading


@dataclass
class Artifact:
    """One artifact type as read: its hits, each a field name -> value map, in order."""

    name: str
    locator: str  # 'Artifact[3]' (XML) or the sheet title (Excel)
    header_locator: str | None = None
    fields: list[str] = field(default_factory=list)  # names as written, first seen first
    hits: list[tuple[str, dict[str, Cell]]] = field(default_factory=list)
    duplicates: list[str] = field(default_factory=list)  # a field name given twice in a hit
    hidden: bool = False

    def add_field(self, name: str) -> None:
        if name not in self.fields:
            self.fields.append(name)


@dataclass
class ReadReport:
    fmt: str  # 'xml' or 'xlsx'
    artifacts: list[Artifact]
    summary: list[tuple[str, str, str]]  # (key, value, locator)
    root: str | None = None
    top_elements: Counter[str] = field(default_factory=Counter)  # the XML root's children
    cells_decoded: int = 0


class _XmlReader:
    """Feeds the XML report through expat; keeps each hit's fragments, nothing else."""

    def __init__(self) -> None:
        self.path: list[str] = []
        self.artifacts: list[Artifact] = []
        self.cur: Artifact | None = None
        self.hit: dict[str, Cell] | None = None
        self.hit_loc = ""
        self.frag: tuple[str, list[str]] | None = None
        self.root: str | None = None
        self.other: Counter[str] = Counter()

    def start(self, name: str, attrs: dict[str, str]) -> None:
        local = name.rpartition(" ")[2]
        if not self.path:
            self.root = local
        self.path.append(local)
        if local == "Artifact":
            n = len(self.artifacts) + 1
            self.cur = Artifact(attrs.get("name", "").strip() or f"(unnamed {n})", f"Artifact[{n}]")
            self.artifacts.append(self.cur)
        elif local == "Hit" and self.cur is not None and self.path[-3:-1] == ["Artifact", "Hits"]:
            self.hit = {}
            self.hit_loc = f"{self.cur.locator}/Hit[{len(self.cur.hits) + 1}]"
        elif local == "Fragment" and self.hit is not None and self.path[-2] == "Hit":
            self.frag = (attrs.get("name", "").strip(), [])
        if len(self.path) == 2:
            self.other[local] += 1

    def end(self, name: str) -> None:
        local = self.path.pop()
        if (
            local == "Fragment"
            and self.frag is not None
            and self.hit is not None
            and self.cur is not None
        ):
            fname, parts = self.frag
            self.frag = None
            self.cur.add_field(fname)
            if fname in self.hit:
                if fname not in self.cur.duplicates:
                    self.cur.duplicates.append(fname)
                return
            value = "".join(parts)
            self.hit[fname] = value if value.strip() else None
        elif (
            local == "Hit"
            and self.hit is not None
            and self.cur is not None
            and self.path[-2:] == ["Artifact", "Hits"]
        ):
            self.cur.hits.append((self.hit_loc, self.hit))
            self.hit = None
        elif local == "Artifact":
            self.cur = None

    def chars(self, data: str) -> None:
        if self.frag is not None:
            self.frag[1].append(data)

    def read(self, stream: IO[bytes]) -> None:
        p = expat.ParserCreate(namespace_separator=" ")
        p.buffer_text = True
        p.StartElementHandler = self.start
        p.EndElementHandler = self.end
        p.CharacterDataHandler = self.chars
        p.StartDoctypeDeclHandler = _refuse_doctype
        p.EntityDeclHandler = _refuse_entity
        p.SetParamEntityParsing(expat.XML_PARAM_ENTITY_PARSING_NEVER)
        try:
            while chunk := stream.read(CHUNK):
                p.Parse(chunk, False)
            p.Parse(b"", True)
        except expat.ExpatError as err:
            raise ValueError(f"the XML report is not well-formed XML ({err})") from err
        if not self.artifacts:
            raise ValueError(
                "the XML file holds no Artifact elements, so it is not a Magnet AXIOM XML report "
                "Baker can read"
            )


def _refuse_doctype(*_a: object) -> None:
    raise ValueError("the XML report declares a document type; Baker does not read such XML")


def _refuse_entity(*_a: object) -> None:
    raise ValueError("the XML report declares an entity; Baker does not read such XML")


class _HashingReader:
    """A read-only stream that SHA-256 hashes what passes through it."""

    def __init__(self, raw: IO[bytes]) -> None:
        self.raw = raw
        self.h = hashlib.sha256()

    def read(self, n: int = -1) -> bytes:
        b = self.raw.read(n)
        self.h.update(b)
        return b


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with open_input(path) as f:
        while chunk := f.read(CHUNK):
            h.update(chunk)
    return h.hexdigest()


def read_xml(path: Path) -> tuple[ReadReport, str]:
    """The XML report and the SHA-256 of exactly the bytes parsed."""
    rd = _XmlReader()
    with open_input(path) as raw:
        hashing = _HashingReader(raw)
        rd.read(hashing)
    rep = ReadReport("xml", rd.artifacts, [], rd.root, rd.other)
    return rep, hashing.h.hexdigest()


_XSTRING_RE = re.compile(r"_x([0-9A-Fa-f]{4})_")


def _known_field(name: object) -> bool:
    if not isinstance(name, str) or not name.strip():
        return False
    b = field_name(name).base
    return b in IGNORED_FIELDS or any(
        b in aliases for f in FIELDS.values() for aliases in f.values()
    )


def read_workbook(data: bytes) -> ReadReport:
    """The Excel report: a case sheet (key, value pairs) and one sheet per artifact type, whose
    header is the first of its top rows holding two or more field names Baker knows (or one,
    when it is the row's only value)."""
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    arts: list[Artifact] = []
    summary: list[tuple[str, str, str]] = []
    decoded = 0
    try:
        for ws in wb.worksheets:
            title = ws.title
            rows: list[tuple[str, list[Cell]]] = []
            for i, cells in enumerate(ws.iter_rows(values_only=True), start=1):
                out: list[Cell] = []
                for c in cells:
                    if isinstance(c, str) and "_x" in c:
                        d = _XSTRING_RE.sub(lambda m: chr(int(m.group(1), 16)), c)
                        decoded += d != c
                        c = d
                    out.append(c)
                rows.append((f"{title}!{i}", out))
            if norm(title) in SUMMARY_SHEETS:
                for loc, cells in rows:
                    vals = [v for v in (text(c) for c in cells) if v is not None]
                    if len(vals) >= 2:
                        summary.append((vals[0], vals[1], loc))
                continue
            art = Artifact(title, title, hidden=getattr(ws, "sheet_state", "visible") != "visible")
            hidx = next(
                (
                    i
                    for i, (_l, cells) in enumerate(rows[:HEADER_SCAN_ROWS])
                    if sum(_known_field(c) for c in cells)
                    >= min(2, sum(text(c) is not None for c in cells) or 2)
                ),
                None,
            )
            if hidx is None:
                art.header_locator = None
                art.hits = [(loc, {}) for loc, cells in rows if any(text(c) for c in cells)]
                arts.append(art)
                continue
            art.header_locator, header = rows[hidx]
            names = [" ".join(str(h).split()) if text(h) is not None else None for h in header]
            for n in names:
                if n is None:
                    continue
                if n in art.fields:
                    if n not in art.duplicates:
                        art.duplicates.append(n)
                else:
                    art.fields.append(n)
            for loc, cells in rows[hidx + 1 :]:
                if not any(text(c) is not None for c in cells):
                    continue
                hit: dict[str, Cell] = {}
                for n, c in zip(names, cells, strict=False):
                    if n is not None and n not in hit:
                        hit[n] = c
                art.hits.append((loc, hit))
            arts.append(art)
    finally:
        wb.close()
    return ReadReport("xlsx", arts, summary, cells_decoded=decoded)


_TIME_HEADER_RE = re.compile(r"date/time\s+-\s+(?:utc|gmt)", re.I)
_PLATFORM_START_RE = re.compile(r"^(?:android|ios)\s")


def looks_like_axiom_workbook(path: Path) -> bool:
    """True when an .xlsx has the marks of an AXIOM export: a time column named with its UTC
    offset ('... Date/Time - UTC+00:00 (M/d/yyyy)'), a 'Recovery method' column, or a sheet
    named as an AXIOM artifact type with its platform ('Android Contacts', 'iOS WhatsApp
    Accounts Information'). Cellebrite reports carry none of these."""
    try:
        with open_input(path) as f:
            wb = load_workbook(io.BytesIO(f.read()), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 - not a workbook: let the Cellebrite reader say why
        return False
    try:
        for ws in wb.worksheets:
            title = norm(ws.title)
            cut = len(ws.title) == EXCEL_SHEET_NAME_LIMIT
            if _PLATFORM_START_RE.match(title) and classify(ws.title, cut)[0] is not None:
                return True
            for cells in ws.iter_rows(max_row=HEADER_SCAN_ROWS, values_only=True):
                names = {norm(c) for c in cells if isinstance(c, str)}
                if any(_TIME_HEADER_RE.search(n) for n in names) or "recovery method" in names:
                    return True
    finally:
        wb.close()
    return False


# ---------------------------------------------------------------- rows


@dataclass
class _HitRow(Row):
    """A Row whose parties are built from separate identifier and name fields."""

    party_lists: dict[str, list[Party]] = field(default_factory=dict)

    def parties(self, f: str) -> list[Party]:
        return list(self.party_lists.get(f, []))


def _parties(ident: Cell, name: Cell) -> list[Party]:
    """Parties from an identifier field and its name field. One name goes with one party; a
    name with no identifier is a name-only party; several names stay unpaired."""
    ps = parse_parties(ident)
    n = text(name)
    if not ps:
        return [Party(None, n, False, None)] if n else []
    if n and len(ps) == 1 and ps[0].name is None:
        return [replace(ps[0], name=n)]
    return ps


@dataclass
class _Mapped:
    kind: RecordKind
    app: str | None
    tc: TableCoverage
    groups: list[tuple[list[Row], TimeContext]]  # rows by the time field they were read from


class _Mapper:
    def __init__(self, cov: Coverage, stated: str | None) -> None:
        self.cov = cov
        self.stated = stated
        self.tables: list[TableCoverage] = []
        self.flag_counts: dict[str, Counter[str]] = {}
        self.unknown: dict[str, list[str]] = {}
        self.recovery_values: Counter[str] = Counter()
        self.device: dict[str, tuple[str, str]] = {}  # key -> (value, locator)
        self.artifact_counts: dict[str, int] = {}
        self.device_fields_not_read: dict[str, list[str]] = {}
        # forwarded and quoted marks by locator (core/ingest/forwarded.py): reason -> locators,
        # and locator -> the original author as written
        self.marks: dict[str, list[str]] = {}
        self.authors: dict[str, str] = {}
        self.copies: list[Row] = []  # embedded copies of the artifact being mapped

    def _resolve(
        self, art: Artifact, kind: str, tc: TableCoverage
    ) -> tuple[dict[str, FieldName], list[FieldName], list[str]]:
        """field key -> the field name used; every time field present, in order of preference
        (a hit with the first one blank takes its time from the next); and notes."""
        aliases = FIELDS[kind]
        names = [field_name(n) for n in art.fields]
        by_base: dict[str, FieldName] = {}
        for fn in names:
            if fn.base in by_base:
                tc.duplicate_columns.append(
                    {"column": fn.written, "field": fn.base, "used": by_base[fn.base].written}
                )
            else:
                by_base[fn.base] = fn
        used: dict[str, FieldName] = {}
        for key, options in aliases.items():
            for a in options:
                if a in by_base:
                    used[key] = by_base[a]
                    break
        known = {a for options in aliases.values() for a in options} | IGNORED_FIELDS
        tc.unknown_columns = [fn.written for fn in names if fn.base not in known]
        notes = []
        ts_fields = [by_base[a] for a in aliases.get("ts", ()) if a in by_base]
        if len(ts_fields) > 1:
            notes.append(
                f"times were read from {ts_fields[0].written!r}; where it is blank, from "
                + ", then ".join(repr(f.written) for f in ts_fields[1:])
                + "."
            )
        for d in art.duplicates:
            tc.duplicate_columns.append({"column": d, "field": norm(d), "used": d})
        return used, ts_fields, notes

    def _flag(self, name: str, tc: TableCoverage, reason: str, locator: str) -> None:
        self.flag_counts.setdefault(name, Counter())[reason] += 1
        _listed(tc.flagged_locators, reason, locator)

    def artifact(self, art: Artifact) -> _Mapped | None:
        truncated = art.locator == art.name and len(art.name) == EXCEL_SHEET_NAME_LIMIT
        kind, app = classify(art.name, truncated)
        self.artifact_counts[art.name] = len(art.hits)
        if kind is None:
            # A sheet with no header Baker knows is counted in rows, any header included.
            count = "rows" if art.locator == art.name and art.header_locator is None else "records"
            self.cov.not_imported.append(
                {"name": art.name, "reason": not_imported_reason(art.name), count: len(art.hits)}
            )
            return None
        if kind == "device":
            known = {a for options in _DEVICE_FIELDS.values() for a in options} | IGNORED_FIELDS
            read = 0
            for loc, hit in art.hits:
                fmap = {field_name(k).base: (k, v) for k, v in hit.items()}
                for key, options in _DEVICE_FIELDS.items():
                    for a in options:
                        if a in fmap and text(fmap[a][1]):
                            read += 1
                            if key not in self.device:
                                self.device[key] = (str(text(fmap[a][1])), loc)
            unknown = [f for f in art.fields if field_name(f).base not in known]
            if unknown:
                self.device_fields_not_read[art.name] = unknown
            if not read:
                self.cov.not_imported.append(
                    {
                        "name": art.name,
                        "reason": "device details with no field Baker knows",
                        "rows" if art.header_locator is None and art.locator == art.name
                        else "records": len(art.hits),
                    }
                )  # fmt: skip
            return None
        if art.header_locator is None and art.locator == art.name:  # an Excel sheet, no header
            self.cov.not_imported.append(
                {
                    "name": art.name,
                    "reason": "no header row with field names Baker knows",
                    "rows": len(art.hits),
                }
            )
            return None
        rkind = RecordKind(kind)
        tc = TableCoverage(name=art.name, kind=kind, header_locator=art.header_locator)
        tc.app_from_title = app
        if art.hidden:
            self.cov.hidden_sheets.append(art.name)
        used, ts_fields, notes = self._resolve(art, kind, tc)
        tc.header_notes = notes
        if truncated:
            tc.header_notes.append(
                "the sheet name is at Excel's 31-character limit and may be cut short; its "
                "kind was read from the start of its last word."
            )
        if rkind is RecordKind.SMS and "chat_id" in used:
            rkind = RecordKind.CHATS  # a thread id is given: group by it, as for chats
            tc.kind = rkind.value
            tc.header_notes.append(
                f"messages were grouped into threads by {used['chat_id'].written!r}."
            )
        if rkind in (RecordKind.CHATS, RecordKind.SMS) and not (
            {"forwarded", "forwarded_from", "quoted_body"} & set(used)
        ):
            tc.header_notes.append(
                "no forwarded or quoted field Baker knows was found, so forwarded or quoted "
                "messages in this table are not marked and may be credited to the forwarder."
            )
        self.tables.append(tc)
        for fn in ts_fields:
            if fn.offset is not None:
                tc.header_offsets[fn.written] = fn.offset
            elif fn.zone is not None:
                tc.header_zone_names[fn.written] = fn.zone
        rows: dict[str, list[Row]] = {fn.written: [] for fn in ts_fields} or {"": []}
        self.copies = []
        n_copies = 0
        for loc, hit in art.hits:
            if all(
                text(hit.get(fn.written)) is None and not isinstance(hit.get(fn.written), datetime)
                for key, fn in used.items()
                if key != "recovery"
            ) and all(text(hit.get(fn.written)) is None for fn in ts_fields):
                tc.rows_skipped["no_values_in_known_columns"] += 1
                _listed(tc.skipped_locators, "no_values_in_known_columns", loc)
                continue
            ts_used = next(
                (fn for fn in ts_fields if _filled(hit.get(fn.written))),
                ts_fields[0] if ts_fields else None,
            )
            row = self._row(art.name, rkind, tc, used, ts_used, loc, hit)
            group = rows[ts_used.written if ts_used else ""]
            group.append(row)
            group.extend(self.copies)  # its embedded copy, after it, in the same group
            n_copies += len(self.copies)
            self.copies = []
        tc.rows_read = len(art.hits)
        # embedded copies are not the export's own hits: counted apart (embedded_messages_imported)
        tc.rows_imported = sum(len(r) for r in rows.values()) - n_copies
        tc.missing_key_fields = [
            k
            for k, f in (("time", "ts"), ("body", "body"), ("from", "from_id"), ("name", "name"))
            if k in {"chats": ["time", "body", "from"], "sms": ["time", "body"],
                     "calls": ["time"], "contacts": ["name"], "accounts": []}[rkind.value]
            and f not in used
            and not (k == "from" and "from_name" in used)
            and not (k == "name" and "first_name" in used)
        ]  # fmt: skip
        tc.absent_fields = sorted(
            f for f in ("ts", "deleted") if f not in used and rkind not in (
                RecordKind.CONTACTS, RecordKind.ACCOUNTS)
        )  # fmt: skip
        by_name = {fn.written: fn for fn in ts_fields}
        groups = []
        for written, group in rows.items():
            own = DateOrder()
            for r in group:
                v = r.get("ts")
                if isinstance(v, str):
                    own.observe(v)
            fn = by_name.get(written)
            offsets = {"ts": fn.offset} if fn is not None and fn.offset is not None else {}
            order = self._order(tc, rkind, own, fn)
            groups.append((group, TimeContext(order, None, offsets, tc)))
        return _Mapped(rkind, app, tc, groups)

    def _order(
        self, tc: TableCoverage, kind: RecordKind, own: DateOrder, ts: FieldName | None
    ) -> DateOrder:
        """The day/month order times are read with: stated at import, else the date format in
        the time field's name, else the table's own values. Never assumed."""
        if kind in (RecordKind.CONTACTS, RecordKind.ACCOUNTS) or ts is None:
            return own
        header = date_format_order(ts.date_format)
        if self.stated is not None:
            eff = DateOrder(int(self.stated == "dmy"), int(self.stated == "mdy"))
            tc.date_order_basis = "stated at import"
            want = self.stated
        elif header in ("dmy", "mdy"):
            eff = DateOrder(int(header == "dmy"), int(header == "mdy"))
            tc.date_order_basis = "field name"
            want = header
        else:
            tc.date_order_basis = "this table" if own.label != "undecided" else "none"
            tc.date_order = own.label
            return own
        if own.label not in ("undecided", want):
            counter = (
                "values_contradicting_stated_date_order"
                if self.stated is not None
                else "values_contradicting_field_date_format"
            )
            self.cov.counters[counter] += own.mdy_votes if want == "dmy" else own.dmy_votes
        tc.date_order = eff.label
        return eff

    def _row(
        self,
        artifact: str,
        kind: RecordKind,
        tc: TableCoverage,
        used: dict[str, FieldName],
        ts: FieldName | None,
        loc: str,
        hit: dict[str, Cell],
    ) -> Row:
        def get(key: str) -> Cell:
            fn = ts if key == "ts" else used.get(key)
            return hit.get(fn.written) if fn is not None else None

        rec = text(get("recovery"))
        if rec is not None:
            self.recovery_values[rec] += 1
            if _CARVED_RE.search(rec):
                self.cov.counters["records_recovered_by_carving"] += 1
                self._flag(artifact, tc, "carved", loc)
        values: dict[str, Cell] = {}
        parties: dict[str, list[Party]] = {}
        attachments: list[str] = []
        if kind in (RecordKind.CHATS, RecordKind.SMS):
            senders = _parties(get("from_id"), get("from_name"))
            to = _parties(get("to_id"), get("to_name"))
            partner = _parties(get("partner_id"), get("partner_name"))
            parties = {"from": senders, "to": to}
            if kind is RecordKind.CHATS:
                parties["participants"] = partner
            else:
                parties["parties"] = partner
            values = {
                "chat_id": get("chat_id"),
                "chat_name": get("chat_name"),
                "body": get("body"),
                "ts": get("ts"),
                "direction": get("direction"),
                "type": get("type"),
                "status": get("status"),
                "deleted": get("deleted"),
                "tag": get("tag"),
            }
            att = text(get("attachment"))
            attachments = [a for a in (att or "").splitlines() if a.strip()]
            if (
                "ts" in used
                and text(values["ts"]) is None
                and not isinstance(values["ts"], datetime)
            ):
                self._flag(artifact, tc, "no_time", loc)
            self._forwarded(artifact, tc, loc, get, values)
        elif kind is RecordKind.CALLS:
            partner = _parties(get("partner_id"), get("partner_name"))
            parties = {
                "from": _parties(get("from_id"), None),
                "to": _parties(get("to_id"), None),
                "parties": partner,
            }
            values = {
                "ts": get("ts"),
                "duration": get("duration"),
                "direction": get("direction"),
                "status": get("status"),
                "deleted": get("deleted"),
            }
            raw = get("duration")
            if text(raw) is not None and _duration(raw) is None:
                self.cov.counters["call_durations_unread"] += 1
                self._flag(artifact, tc, "duration_unread", loc)
            values["duration"] = _duration(raw)
            if (
                "ts" in used
                and text(values["ts"]) is None
                and not isinstance(values["ts"], datetime)
            ):
                self._flag(artifact, tc, "no_time", loc)
        elif kind is RecordKind.CONTACTS:
            name = text(get("name")) or " ".join(
                x for x in (text(get("first_name")), text(get("last_name"))) if x
            )
            values = {
                "name": name or None,
                "phone": get("phone"),
                "email": get("email"),
                "username": get("username"),
            }
        else:  # accounts
            values = {"username": get("username"), "name": get("name")}
        cells: list[Cell] = [*values.values(), *attachments]
        return _HitRow(
            loc,
            cells,
            {k: i for i, k in enumerate(values)},
            list(range(len(values), len(cells))),
            party_lists=parties,
        )

    def _forwarded(
        self,
        artifact: str,
        tc: TableCoverage,
        loc: str,
        get: Callable[[str], Cell],
        values: dict[str, Cell],
    ) -> None:
        """Mark a forwarded or quoting hit the way the UFDR importer does (core/ingest/forwarded.py)
        and queue its quoted original, one level deep, as a copy at loc + '/Embedded[1]'."""
        is_fwd = _forwarded_value(get("forwarded"))
        quoted = get("quoted_body")
        copy_senders = _parties(get("quoted_from_id"), get("quoted_from_name"))
        has_copy = text(quoted) is not None
        if is_fwd:
            self.cov.counters["messages_marked_forwarded"] += 1
            self.marks.setdefault("forwarded", []).append(loc)
            # the original author as written: the stated field, else the copy's one sender
            author = text(get("forwarded_from"))
            if author is None and has_copy:
                author = _sole(get("quoted_from_id"), get("quoted_from_name"), copy_senders)
            if author is not None:
                self.authors[loc] = author
            self._flag(artifact, tc, "forwarded", loc)
        if not has_copy:
            return
        self.marks.setdefault("embedded", []).append(loc)
        self.cov.counters["messages_with_embedded_message_imported_separately"] += 1
        self._flag(artifact, tc, "embedded_message_imported_separately", loc)
        eloc = f"{loc}{forwarded.COPY_STEP}1]"
        self.cov.counters["embedded_messages_imported"] += 1
        self.marks.setdefault("embedded_copy", []).append(eloc)
        author = _sole(get("quoted_from_id"), get("quoted_from_name"), copy_senders)
        if author is not None:
            self.authors[eloc] = author
        self._flag(artifact, tc, "embedded_original_imported", eloc)
        # Its body verbatim, its own sender as written, the parent's thread (fixed up in _write);
        # no time (the original's time is not when this phone held it) and no stated direction.
        cv: dict[str, Cell] = {
            "chat_id": values["chat_id"],
            "chat_name": values["chat_name"],
            "body": quoted,
            "ts": None,
        }
        self.copies.append(
            _HitRow(
                eloc,
                list(cv.values()),
                {k: i for i, k in enumerate(cv)},
                [],
                party_lists={"from": copy_senders},
            )
        )


_FALSE_VALUES = {"no", "false", "0", "n", "not forwarded"}


def _forwarded_value(v: Cell) -> bool:
    """A forwarded flag field says the message was forwarded. Blank or a stated no is not;
    any other value (Yes, True, 1, Forwarded, or one Baker does not know) fails closed."""
    if isinstance(v, bool):
        return v
    t = text(v)
    return t is not None and t.lower() not in _FALSE_VALUES


def _sole(ident: Cell, name: Cell, parties: list[Party]) -> str | None:
    """The one sender as written (identifier, else name); None unless exactly one is named."""
    if len(parties) != 1:
        return None
    return text(ident) or text(name)


_TIMESPAN_RE = re.compile(r"^(?:(\d+)[.:])?(\d{1,2}):(\d{2}):(\d{2})(?:\.\d+)?$")


def _filled(v: Cell) -> bool:
    return isinstance(v, datetime) or text(v) is not None


def _duration(raw: Cell) -> int | None:
    """Seconds from '00:03:38', '1.02:03:04', '218' or a number; None when unread."""
    if isinstance(raw, int | float) and not isinstance(raw, bool):
        return int(raw)
    t = text(raw)
    if t is None:
        return None
    if re.fullmatch(r"\d+", t):
        return int(t)
    m = _TIMESPAN_RE.match(t)
    if m:
        d, h, mi, s = m.groups()
        return int(d or 0) * 86400 + int(h) * 3600 + int(mi) * 60 + int(s)
    m = re.fullmatch(r"(\d{1,2}):(\d{2})", t)
    if m:
        return int(m.group(1)) * 60 + int(m.group(2))
    return None


# ---------------------------------------------------------------- the importer


class AxiomImporter:
    """Implements core.contracts.EvidenceImporter for Magnet AXIOM report exports: the XML
    report (.xml) and the Excel report (.xlsx). Format support is UNVERIFIED (module docstring);
    each import records that."""

    name = IMPORTER_NAME

    def __init__(
        self,
        source_id: str | None = None,
        now: datetime | None = None,
        date_order: str | None = None,
        expected_sha256: str | None = None,
    ) -> None:
        """`date_order` ('mdy' or 'dmy'): the day/month order a person states. Without it, the
        order comes from the date format in each time field's name, then from values that
        decide it; otherwise ambiguous dates are reported unread."""
        if date_order not in (None, "mdy", "dmy"):
            raise ValueError(f"date_order must be 'mdy' or 'dmy', not {date_order!r}")
        self.source_id = source_id
        self.now = now
        self.date_order = date_order
        self.expected_sha256 = expected_sha256  # refuse other bytes (source_file.require_sha256)

    def import_source(self, path: Path, conn: sqlite3.Connection) -> Source:
        if path.suffix.lower() == ".xml":
            if self.expected_sha256 is not None:  # refuse other bytes before parsing them
                require_sha256(path.name, self.expected_sha256, sha256_file(path))
            report, sha = read_xml(path)
        else:
            with open_input(path) as f:
                data = f.read()
            sha = hashlib.sha256(data).hexdigest()
            require_sha256(path.name, self.expected_sha256, sha)
            try:
                report = read_workbook(data)
            except Exception as err:  # noqa: BLE001 - say what the file is not
                raise ValueError(f"{path.name} is not an Excel workbook Baker can read") from err
        require_sha256(path.name, self.expected_sha256, sha)  # the XML parsed is the XML checked
        return self._write(path, sha, conn, report)

    def _write(self, path: Path, sha: str, conn: sqlite3.Connection, rep: ReadReport) -> Source:
        src = self.source_id or _source_id(path, sha)
        if conn.execute("SELECT 1 FROM sources WHERE id = ?", (src,)).fetchone():
            raise ValueError(f"source {src} is already in this case database")
        cov = Coverage()
        cov.counters["cells_with_excel_escapes_decoded"] = rep.cells_decoded
        mp = _Mapper(cov, self.date_order)
        mapped = [m for m in (mp.artifact(a) for a in rep.artifacts) if m is not None]
        cov.summary_found = bool(rep.summary) or bool(mp.device)
        cov.tables = mp.tables

        summ, summ_names = _summary(rep.summary)
        extracted_utc = None
        if "extraction_start" in summ:
            fn = summ_names["extraction_start"]
            named = date_format_order(fn.date_format)
            stated = self.date_order or (named if named in ("dmy", "mdy") else None)
            order = DateOrder(int(stated == "dmy"), int(stated == "mdy")) if stated else DateOrder()
            extracted_utc = parse_time(summ["extraction_start"][0], order, fn.offset).ts.utc
            if extracted_utc is None:
                cov.counters["extraction_time_unread"] += 1
        dev_label = _device_label(mp.device, summ)
        source = Source(
            id=src,
            kind=SourceKind.AXIOM,
            fidelity=Fidelity.CURATED_REPORT,
            extraction_type=extraction_type(summ.get("extraction_type", (None, ""))[0]),
            file_name=path.name,
            sha256=sha,
            tool_name="Magnet AXIOM (report export)",
            tool_version=summ.get("tool_version", (None, ""))[0],
            extracted_at_utc=extracted_utc,
            imported_at_utc=self.now or datetime.now(UTC),
        )
        dev = f"dev:{src}"
        dev_loc = (
            mp.device["device"][1]
            if "device" in mp.device
            else summ["device"][1]
            if "device" in summ
            else ""
        )
        w = _Writer(conn, src, dev, cov)
        order = [RecordKind.ACCOUNTS, RecordKind.CONTACTS, RecordKind.CHATS, RecordKind.SMS,
                 RecordKind.CALLS]  # fmt: skip
        with conn:
            conn.execute(
                "INSERT INTO sources (id, kind, fidelity, extraction_type, file_name, sha256,"
                " tool_name, tool_version, extracted_at_utc, imported_at_utc)"
                " VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                (
                    source.id,
                    source.kind.value,
                    source.fidelity.value,
                    source.extraction_type.value,
                    source.file_name,
                    source.sha256,
                    source.tool_name,
                    source.tool_version,
                    _iso(source.extracted_at_utc),
                    _iso(source.imported_at_utc),
                ),
            )
            conn.execute(
                "INSERT INTO devices (id, source_id, locator, label, model, os_version, timezone)"
                " VALUES (?, ?, ?, ?, ?, ?, ?)",
                (
                    dev,
                    src,
                    dev_loc,
                    dev_label or f"Device in {path.name}",
                    dev_label,
                    (mp.device.get("os") or summ.get("os") or (None, ""))[0],
                    mp.device.get("time_zone", (None, ""))[0],
                ),
            )
            for kind in order:
                for m in mapped:
                    if m.kind is not kind:
                        continue
                    for rows, tctx in m.groups:
                        if kind is RecordKind.ACCOUNTS:
                            _accounts(w, rows, m.tc, m.app)
                        elif kind is RecordKind.CONTACTS:
                            _contacts(w, rows)
                        elif kind is RecordKind.CALLS:
                            _calls(w, rows, tctx, m.app)
                        else:
                            _messages(w, kind, rows, tctx, m.app)
            _attach_copies(w, mp)
            w.flush()
            forwarded.mark(conn, cov.forwarded_records, cov.forwarded_from)
            payload = coverage_payload(
                source,
                IMPORTER_NAME,
                None,
                None,
                cov,
                sorted({m.kind for m in mapped}),
                None,
                self.date_order,
            )
            payload["tables"] = [_table_dict(tc, mp) for tc in cov.tables]
            payload.update(
                importer_version=IMPORTER_VERSION,
                fidelity_reason=FIDELITY_REASON,
                format_support=FORMAT_SUPPORT,
                axiom={
                    "format": rep.fmt,
                    "root_element": rep.root,
                    "top_elements": dict(sorted(rep.top_elements.items())),
                    "artifacts": dict(mp.artifact_counts),
                    "recovery_method_values": dict(sorted(mp.recovery_values.items())),
                    "device_read": {k: v[1] for k, v in sorted(mp.device.items())},
                    "device_fields_not_read": mp.device_fields_not_read,
                    "summary_read": {k: v[1] for k, v in sorted(summ.items())},
                },
            )
            conn.execute(
                "INSERT INTO source_coverage (source_id, importer, importer_version, payload_json)"
                " VALUES (?, ?, ?, ?)",
                (source.id, IMPORTER_NAME, IMPORTER_VERSION, json.dumps(payload, sort_keys=True)),
            )
            audit_log.append(conn, source.imported_at_utc, "system", "import", payload)
            if rep.fmt == "xml" and sha256_file(path) != sha:
                raise ValueError(
                    f"{path.name} changed while Baker read it; nothing was imported. Nothing "
                    "should change a copied input, so check the case folder"
                )
        return source


def _summary(
    pairs: list[tuple[str, str, str]],
) -> tuple[dict[str, tuple[str, str]], dict[str, FieldName]]:
    """key -> (value, locator) for the case sheet values Baker reads, and key -> the name as
    written. A key may carry a zone and date format, as a time field's name does
    ('Acquisition Date/Time - UTC+00:00 (M/d/yyyy)')."""
    out: dict[str, tuple[str, str]] = {}
    names: dict[str, FieldName] = {}
    parsed = [(field_name(k.rstrip(": ")), v, loc) for k, v, loc in pairs]
    for key, aliases in SUMMARY_KEYS.items():
        for a in aliases:
            hit = next(((fn, v, loc) for fn, v, loc in parsed if fn.base == a and v.strip()), None)
            if hit:
                out[key] = (hit[1].strip(), hit[2])
                names[key] = hit[0]
                break
    return out, names


def _device_label(
    device: dict[str, tuple[str, str]], summ: dict[str, tuple[str, str]]
) -> str | None:
    model = (device.get("device") or summ.get("device") or (None, ""))[0]
    make = device.get("manufacturer", (None, ""))[0]
    if model and make and not norm(model).startswith(norm(make)):
        return f"{make} {model}"
    return model


def _attach_copies(w: _Writer, mp: _Mapper) -> None:
    """Before the writer flushes: put each embedded copy in its parent's thread (a thread made
    only for a copy is dropped), and list forwarded marks by record id for coverage and for
    forwarded.mark, as the UFDR importer does."""
    loc_i, tid_i = 2, 3  # positions in the writer's message tuple
    by_loc = {m[loc_i]: m[tid_i] for m in w.messages}
    for i, m in enumerate(w.messages):
        loc = m[loc_i]
        if forwarded.COPY_STEP in loc:
            parent = by_loc.get(loc.rsplit(forwarded.COPY_STEP, 1)[0])
            if parent is not None:
                w.messages[i] = (*m[:tid_i], parent, *m[tid_i + 1 :])
    used = {m[tid_i] for m in w.messages}
    for tid in [t for t, v in w.threads.items() if t not in used and forwarded.COPY_STEP in v[0]]:
        del w.threads[tid]
    for reason, locs in mp.marks.items():
        w.cov.forwarded_records.setdefault(reason, []).extend(f"msg:{w.src}:{loc}" for loc in locs)
    for loc, author in mp.authors.items():
        w.cov.forwarded_from[f"msg:{w.src}:{loc}"] = author


def _table_dict(tc: TableCoverage, mp: _Mapper) -> dict[str, object]:
    d = tc.as_dict()
    d["flag_counts"] = dict(mp.flag_counts.get(tc.name, {}))
    return d


def _source_id(path: Path, sha: str) -> str:
    """The file stem when it is a plain token, else `src_<sha12>` (as the other importers)."""
    if re.fullmatch(r"[A-Za-z0-9_\-]{1,64}", path.stem):
        return path.stem
    return f"src_{sha[:12]}"


def _iso(dt: datetime | None) -> str | None:
    return dt.astimezone(UTC).isoformat().replace("+00:00", "Z") if dt else None
