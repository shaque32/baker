"""Cellebrite-style Excel report layouts for the variant set. All synthetic.

`sheets()` renders a phone in the case01 layout; each variant changes one thing about it.
The workbook is written with openpyxl so merged cells, real date cells and hidden sheets can be
expressed. Variants are rebuilt in a temporary folder by the tests, never committed.
"""

from __future__ import annotations

from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

from eval.synthetic.variants.base import (
    DEVICE_OFFSET,
    PHONE,
    Call,
    Msg,
    Phone,
    hms,
    us12,
)

Cell = str | int | float | datetime | None
TimeFmt = Callable[[datetime, int], Cell]
FIXED = datetime(2026, 3, 20, 12, 0, 0)  # workbook properties, so files do not carry the clock


@dataclass
class SheetSpec:
    title: str
    rows: list[list[Cell]]
    header_rows: int = 1  # rows from the top up to and including the header
    merges: list[str] = field(default_factory=list)
    hidden: bool = False

    def data_rows(self) -> int:
        """Rows under the header holding anything (what an importer must account for)."""
        return sum(
            1
            for r in self.rows[self.header_rows :]
            if any(c is not None and str(c).strip() for c in r)
        )


def write(path: Path, sheets: list[SheetSpec]) -> dict[str, int]:
    """Write the workbook; returns the data-row count per sheet."""
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.created = FIXED
    wb.properties.modified = FIXED
    wb.properties.creator = "baker synthetic variant"
    for spec in sheets:
        ws = wb.create_sheet(spec.title)
        for r in spec.rows:
            ws.append(list(r))
        for ref in spec.merges:
            ws.merge_cells(ref)
        if spec.hidden:
            ws.sheet_state = "hidden"
    wb.save(path)
    return {s.title: s.data_rows() for s in sheets}


# ---------------------------------------------------------------- the default layout


def summary(time_zone: str | None = "UTC-4") -> SheetSpec:
    rows: list[list[Cell]] = [
        ["Extraction Summary"],
        ["Case number", "SYN-VARIANT-0001"],
        ["Device", "SyntheticPhone V2"],
        ["OS version", "17.2"],
        ["Extraction type", "Advanced Logical"],
        ["Extraction start date/time", "3/18/2026 10:00:00 AM(UTC-4)"],
        ["UFED Physical Analyzer version", "7.70.0.1"],
    ]
    if time_zone is not None:
        rows.append(["Time zone", time_zone])
    return SheetSpec("Summary", rows, header_rows=0)


def chat_cols(phone: Phone, ts: TimeFmt) -> list[tuple[str, Callable[[int, Msg], Cell]]]:
    def chat(m: Msg):
        return phone.chat(m.chat)

    return [
        ("#", lambda i, m: i),
        ("Chat #", lambda i, m: m.chat),
        ("Name", lambda i, m: chat(m).name),  # type: ignore[union-attr]
        ("Source", lambda i, m: chat(m).app),  # type: ignore[union-attr]
        ("Participants", lambda i, m: "\n".join(w.cell() for w in chat(m).members)),  # type: ignore[union-attr]
        ("From", lambda i, m: m.sender.cell()),
        ("To", lambda i, m: "\n".join(w.cell() for w in m.to)),
        ("Body", lambda i, m: m.body or None),
        ("Timestamp", lambda i, m: ts(m.utc, DEVICE_OFFSET)),
        ("Direction", lambda i, m: "Outgoing" if m.outgoing else "Incoming"),
        ("Deleted", lambda i, m: "Intact"),
        ("Attachment #1", lambda i, m: m.attachments[0] if m.attachments else None),
    ]


def sms_cols(ts: TimeFmt) -> list[tuple[str, Callable[[int, Msg], Cell]]]:
    return [
        ("#", lambda i, m: i),
        ("From", lambda i, m: m.sender.cell()),
        ("To", lambda i, m: "\n".join(w.cell() for w in m.to)),
        ("Body", lambda i, m: m.body),
        ("Timestamp", lambda i, m: ts(m.utc, DEVICE_OFFSET)),
        ("Type", lambda i, m: "Sent" if m.outgoing else "Received"),
        ("Deleted", lambda i, m: "Intact"),
    ]


def call_cols(ts: TimeFmt) -> list[tuple[str, Callable[[int, Call], Cell]]]:
    return [
        ("#", lambda i, c: i),
        ("Source", lambda i, c: c.app),
        ("Type", lambda i, c: c.direction),
        ("Parties", lambda i, c: c.party.cell()),
        ("Timestamp", lambda i, c: ts(c.utc, DEVICE_OFFSET)),
        ("Duration", lambda i, c: hms(c.duration_s)),
        ("Deleted", lambda i, c: "Intact"),
    ]


def table(
    title: str,
    cols: list[tuple[str, Callable]],
    records: list,
    *,
    rename: dict[str, str] | None = None,
    order: list[str] | None = None,
    drop: tuple[str, ...] = (),
    extra: dict[str, Callable] | None = None,
    title_rows: list[list[Cell]] | None = None,
) -> SheetSpec:
    """One data sheet: optional rows above the header, then the header, then one row per
    record. `order` lists headers (before renaming) in the order they appear."""
    cols = [c for c in cols if c[0] not in drop] + list((extra or {}).items())
    if order is not None:
        rank = {h: k for k, h in enumerate(order)}
        cols = sorted(cols, key=lambda c: rank.get(c[0], len(rank)))
    names = [(rename or {}).get(h, h) for h, _f in cols]
    top = [list(r) for r in (title_rows if title_rows is not None else [[title]])]
    rows = [*top, names]
    for i, rec in enumerate(records, start=1):
        rows.append([f(i, rec) for _h, f in cols])
    return SheetSpec(title, rows, header_rows=len(top) + 1)


def sheets(
    phone: Phone = PHONE,
    ts: TimeFmt = us12,
    *,
    time_zone: str | None = "UTC-4",
    chats_title: str = "Chats",
    **chat_kw: object,
) -> list[SheetSpec]:
    """The default layout: Summary, User Accounts, Contacts, Chats, SMS Messages, Call Log."""
    chats = [m for m in phone.messages if m.chat is not None]
    sms = [m for m in phone.messages if m.chat is None]
    return [
        summary(time_zone),
        SheetSpec(
            "User Accounts",
            [["User Accounts"], ["#", "Source", "Username", "Name"]]
            + [[i, a, u, n] for i, (a, u, n) in enumerate(phone.accounts, start=1)],
            header_rows=2,
        ),
        SheetSpec(
            "Contacts",
            [["Contacts"], ["#", "Name", "Entries", "Source"]]
            + [
                [i, c.name, "\n".join(c.entries), "Phone"]
                for i, c in enumerate(phone.contacts, start=1)
            ],
            header_rows=2,
        ),
        table(chats_title, chat_cols(phone, ts), chats, **chat_kw),  # type: ignore[arg-type]
        table("SMS Messages", sms_cols(ts), sms),
        table("Call Log", call_cols(ts), list(phone.calls)),
    ]


def find(specs: list[SheetSpec], title: str) -> SheetSpec:
    return next(s for s in specs if s.title == title)
